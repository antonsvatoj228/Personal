import openpyxl
from openpyxl.utils import get_column_letter
wb = openpyxl.load_workbook('price.xlsx', data_only=True)
km = wb['Керування прайсом']; dov = wb['Довідник категорій']
G6 = km['G6'].value; G7 = km['G7'].value; G8 = km['G8'].value
tariff = {km.cell(row=r,column=6).value: km.cell(row=r,column=7).value for r in range(13,20)}
# price list columns
plcol = {}
for c in range(9,68,2):
    n = km.cell(row=2,column=c).value
    if n: plcol.setdefault(n,c)
basecats = {km.cell(row=r,column=8).value: r for r in range(12,16)}
matrows  = {km.cell(row=r,column=8).value: r for r in range(20,38)}
spar     = {km.cell(row=r,column=8).value: r for r in range(51,53)}
pkgr     = {km.cell(row=r,column=8).value: r for r in range(42,46)}
cat = {}
for r in range(5,23):
    name=dov.cell(row=r,column=1).value
    cat[name]=dict(book=dov.cell(row=r,column=2).value, grp=dov.cell(row=r,column=3).value,
                   base=dov.cell(row=r,column=4).value, mn=dov.cell(row=r,column=5).value,
                   incl=dov.cell(row=r,column=6).value, mx=dov.cell(row=r,column=7).value,
                   role=dov.cell(row=r,column=8).value, spa=dov.cell(row=r,column=9).value)
def val(r,c):
    v=km.cell(row=r,column=c).value
    return 0 if v in (None,'') else v
def rnd(x): return round(x/G7,0)*G7
def price_sheet(catn, n, pl, plan, day, promo):
    c=plcol[pl]; info=cat[catn]
    base=val(basecats[info['base']], c)
    mult=val(matrows[catn], c); wm=val(matrows[catn], c+1)
    if n > info['mx']: return None
    if base==0 or mult==0: return None
    if day=='Вихідні' and wm==0: return None
    pkg=0
    if plan in ('BB','BBSPA','ROSPABB'):
        key = 'BB|'+day if plan=='ROSPABB' else plan+'|'+day
        pkg=val(pkgr[key], c)
        if pkg==0: return None
    spa=0
    if plan in ('ROSPA','ROSPABB'):
        spa=val(spar[info['spa']], c)
        if spa==0: return None
    disc = G8 if plan=='BBSPA' else tariff.get(promo,0)
    acc = base*mult*(wm if day=='Вихідні' else 1) + max(0,n-2)*G6
    return rnd(acc*(1-disc) + n*pkg + spa)
def price_calc(catn, adults, kids, pl, plan, day, promo):
    c=plcol[pl]; info=cat[catn]; tot=adults+kids
    base=val(basecats[info['base']], c); mult=val(matrows[catn], c); wm=val(matrows[catn], c+1)
    f11=base*mult; f12=mult; f13=wm
    if tot<info['mn'] or tot>info['mx'] or f11==0 or f12==0 or (day=='Вихідні' and f13==0): return None
    if plan in ('RO','ROSPA'): pa=pk=0
    else:
        key='BB|'+day if plan=='ROSPABB' else plan+'|'+day
        pa=val(pkgr[key],c); pk=val(pkgr[key],c+1)
        if (adults>0 and pa==0) or (kids>0 and pk==0): return None
    spa=0
    if plan in ('ROSPA','ROSPABB'):
        spa=val(spar[info['spa']],c)
        if spa==0: return None
    f14=max(0, tot - info['incl'])
    f15=rnd(f11*(f13 if day=='Вихідні' else 1) + f14*G6)
    f16=G8 if plan=='BBSPA' else tariff.get(promo,0)
    f17=rnd(f15*(1-f16))
    f20=adults*pa+kids*pk+spa
    return rnd(f17+f20)
if __name__=='__main__':
    # 1) validate model against cached Прайс values
    pr=wb['Прайс']
    B3=pr['B3'].value; E3=pr['E3'].value; H3=pr['H3'].value; K3=pr['K3'].value
    print('Active:',B3,'|',E3,'|',H3,'|',K3)
    bad=0; checked=0
    for r in list(range(8,15))+list(range(16,27)):
        catn=pr.cell(row=r,column=1).value
        for col in range(4,12):
            n=pr.cell(row=7,column=col).value
            cached=pr.cell(row=r,column=col).value
            mine=price_sheet(catn,int(n),B3,H3,K3,E3)
            cached = None if cached in (None,'') else cached
            mine = None if mine is None else int(mine)
            checked+=1
            if cached!=mine:
                bad+=1; print(f'  MISMATCH {get_column_letter(col)}{r} {catn} n={int(n)}: file={cached} model={mine}')
    print(f'model vs file: {checked-bad}/{checked} match')
