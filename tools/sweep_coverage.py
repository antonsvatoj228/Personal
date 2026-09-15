from model import *
pls=[n for n in plcol]
plans=['RO','BB','BBSPA','ROSPA','ROSPABB']; days=['Будні','Вихідні']
promos=[t for t in tariff if t]
print('TARIFFS:',tariff,'G6',G6,'G7',G7,'G8',G8)
print('PRICELISTS:',pls)
print()
print('=== 1. Прайс vs Калькулятор (kids=0), all pricelists/plans/days/promos ===')
diff=0; tot=0; examples=[]
for pl in pls:
  for plan in plans:
    for day in days:
      for promo in promos:
        for catn,info in cat.items():
          for n in range(1,int(info['mx'])+1):
            a=price_sheet(catn,n,pl,plan,day,promo)
            b=price_calc(catn,n,0,pl,plan,day,promo)
            tot+=1
            if a!=b:
              diff+=1
              if len(examples)<12: examples.append((pl,plan,day,promo,catn,n,a,b))
print(f'  compared {tot}, differ {diff}')
for e in examples: print('   ',e)
print()
print('=== 2. Люкс-hole: Апартаменти Люкс 25 м², 3 гостей ===')
for plan in plans:
  for day in days:
    print(f'   Прайс(F18)=<порожньо, формулу видалено>  Калькулятор({plan},{day}) =',
          price_calc('Апартаменти Люкс 25 м²',3,0,'Осінь 2026 (00-10%)',plan,day,'Rack Rate'))
print()
print('=== 3. Матриця/Базові ціни vs Прайс (RO, 1-2 ос.) для активного прайсу ===')
import openpyxl
wb2=openpyxl.load_workbook('price.xlsx',data_only=True)
mx=wb2['Матриця категорій']
pl='Осінь 2026 (00-10%)'
for r in range(7,25):
    catn=mx.cell(row=r,column=1).value
    e=mx.cell(row=r,column=5).value; f=mx.cell(row=r,column=6).value
    p_wd=price_sheet(catn,2,pl,'RO','Будні','Rack Rate')
    p_we=price_sheet(catn,2,pl,'RO','Вихідні','Rack Rate')
    flag='' if (e==p_wd and f==p_we) else '   <<< РОЗБІЖНІСТЬ'
    print(f'   {catn[:44]:46} матриця={e}/{f}   прайс={p_wd}/{p_we}{flag}')
print()
print('=== 4. Покриття: скільки комбінацій дають порожньо ===')
from collections import Counter
c=Counter()
for pl in pls:
  for plan in plans:
    for day in days:
      n_ok=sum(1 for catn,info in cat.items() for n in range(1,int(info['mx'])+1)
               if price_sheet(catn,n,pl,plan,day,'Rack Rate') is not None)
      c[(pl,plan,day)]=n_ok
tot_cells=sum(int(i['mx']) for i in cat.values())
for k,v in c.items():
    if v==0: print(f'   ПУСТО: {k}')
print(f'   (всього комірок на комбінацію: {tot_cells})')
