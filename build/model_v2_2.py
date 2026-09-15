"""Незалежна модель НОВОЇ логіки округлення (для звірки з файлом)."""
import openpyxl
from decimal import Decimal, ROUND_HALF_UP

def xlround(x, step):
    """ROUND() Excel: половина завжди вгору (Python round() округлює до парного)."""
    q = Decimal(repr(x / step)).quantize(Decimal(1), rounding=ROUND_HALF_UP)
    return float(q) * step
def load(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    km, dov = wb['Керування прайсом'], wb['Довідник категорій']
    G = {a: km[a].value for a in ('G6','G7','G8')}
    tariff = {km.cell(row=r,column=6).value: km.cell(row=r,column=7).value
              for r in range(14,22) if km.cell(row=r,column=6).value}
    plcol = {}
    for c in range(9, 129, 2):
        n = km.cell(row=2,column=c).value
        if n: plcol.setdefault(n, c)
    basecats = {km.cell(row=r,column=8).value: r for r in range(12,16)}
    matrows  = {km.cell(row=r,column=8).value: r for r in range(33,51)}
    pkgr     = {km.cell(row=r,column=8).value: r for r in range(54,58)}
    spar     = {km.cell(row=r,column=8).value: r for r in range(61,63)}
    swr      = {'BB':66,'BBSPA':67,'SPA':68}
    cat = {}
    for r in range(5,23):
        nm = dov.cell(row=r,column=1).value
        if not nm: continue
        cat[nm] = dict(base=dov.cell(row=r,column=4).value, mn=dov.cell(row=r,column=5).value,
                       incl=dov.cell(row=r,column=6).value, mx=dov.cell(row=r,column=7).value,
                       spa=dov.cell(row=r,column=9).value)
    def val(r,c):
        v = km.cell(row=r,column=c).value
        return 0 if v in (None,'') else v
    return dict(km=km, G=G, tariff=tariff, plcol=plcol, basecats=basecats, matrows=matrows,
                pkgr=pkgr, spar=spar, swr=swr, cat=cat, val=val)

def price(M, catn, n, pl, pkg, day, tar, adults=None, kids=0):
    if adults is None: adults, kids = n, 0
    G, val, cat = M['G'], M['val'], M['cat']
    STEP, GS = G['G7'], G['G6']
    if pl not in M['plcol'] or catn not in cat: return None
    c, info = M['plcol'][pl], cat[catn]
    base = val(M['basecats'][info['base']], c)
    mult = val(M['matrows'][catn], c); wm = val(M['matrows'][catn], c+1)
    if n > info['mx'] or base == 0 or mult == 0: return None
    if day == 'Вихідні' and wm == 0: return None
    need_pkg = pkg in ('BB','BBSPA','ROSPABB'); need_spa = pkg in ('ROSPA','ROSPABB')
    pa = pk = 0
    if need_pkg:
        key = ('BB|' if pkg == 'ROSPABB' else pkg+'|') + day
        pa = val(M['pkgr'][key], c); pk = val(M['pkgr'][key], c+1)
        if pa == 0: return None
    spa = 0
    if need_spa:
        spa = val(M['spar'][info['spa']], c)
        if spa == 0: return None
    disc = G['G8'] if pkg == 'BBSPA' else M['tariff'].get(tar, 0)
    sw_pkg = val(M['swr']['BBSPA' if pkg == 'BBSPA' else 'BB'], c)
    sw_spa = val(M['swr']['SPA'], c)
    acc = base * mult * (wm if day == 'Вихідні' else 1)
    return (xlround(acc * (1 - disc), STEP)
            + max(0, n - info['incl']) * GS * (1 - disc)
            + (adults * pa + kids * pk) * (1 - disc * sw_pkg)
            + spa * (1 - disc * sw_spa))
