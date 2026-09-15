import openpyxl, sys
sys.path.insert(0,'.')
from model import *
from pls import json_pls
v2 = openpyxl.load_workbook('v2.xlsx', data_only=True)
orig = openpyxl.load_workbook('price.xlsx', data_only=True)

cat['Апартаменти Люкс 25 м²']['mx'] = 2.0        # єдина узгоджена зміна даних
P = v2['Прайс']; O = orig['Прайс']
PL, TAR, PKG = P['B3'].value, P['B4'].value, P['B5'].value
print(f'v2 «Прайс»: {PL} | {TAR} | {PKG}')
print(f'оригінал  : {O["B3"].value} | {O["E3"].value} | {O["H3"].value} | {O["K3"].value}\n')

ok=bad=0; issues=[]
for i in range(18):
    r = 9+i; name = P.cell(row=r,column=1).value
    for j in range(8):
        n = j+1
        got_b = P.cell(row=r, column=4+j).value
        got_v = P.cell(row=r, column=13+j).value
        exp_b = price_sheet(name, n, PL, PKG, 'Будні',   TAR)
        exp_v = price_sheet(name, n, PL, PKG, 'Вихідні', TAR)
        for got,exp,lab in ((got_b,exp_b,'будні'),(got_v,exp_v,'вих.')):
            got = None if got in (None,'') else round(got)
            exp = None if exp is None else round(exp)
            if got == exp: ok+=1
            else: bad+=1; issues.append(f'  {name[:32]:34} {n} ос. {lab}: файл={got} модель={exp}')
print(f'1) Прайс v2 проти незалежної моделі: {ok} збігів, {bad} розбіжностей')
for x in issues[:10]: print(x)

# 2) будні v2 == оригінальний Прайс (той самий набір параметрів)
same=diff=0; dl=[]
orows = list(range(8,15))+list(range(16,27))
for i,orow in enumerate(orows):
    r = 9+i; name = P.cell(row=r,column=1).value
    assert name == O.cell(row=orow,column=1).value, (name, O.cell(row=orow,column=1).value)
    for j in range(8):
        a = O.cell(row=orow, column=4+j).value
        b = P.cell(row=r,    column=4+j).value
        a = None if a in (None,'') else round(a); b = None if b in (None,'') else round(b)
        if a==b: same+=1
        else: diff+=1; dl.append(f'  {name[:34]:36} {j+1} ос.: оригінал={a} → v2={b}')
print(f'\n2) Будні v2 проти ОРИГІНАЛУ: {same} однакових, {diff} змінених')
for x in dl: print(x)

# 3) Калькулятор == Прайс (лише дорослі)
C = v2['Калькулятор']
print(f'\n3) Калькулятор: {C["B8"].value}, {C["B10"].value} дор. + {C["B11"].value} діт., {C["B9"].value}, {C["B12"].value}')
print(f'   ЦІНА ЗА НІЧ = {C["F23"].value}   очікуємо (модель) = '
      f'{price_sheet(C["B8"].value, int(C["B10"].value)+int(C["B11"].value), C["B6"].value, C["B12"].value, C["B9"].value, C["B7"].value)}')

# 4) дані всіх прайс-листів збережені
K = v2['Керування прайсом']; KO = orig['Керування прайсом']
miss=[]
for i,pl in enumerate(json_pls()):
    c2, co = 9+i*2, 9+i*2
    for rr2,rro in [(12,12),(13,13),(14,14),(15,15)]:
        a,b = KO.cell(row=rro,column=co).value, K.cell(row=rr2,column=c2).value
        if a!=b: miss.append(f'  база {pl} р.{rro}: {a} → {b}')
    for k in range(18):
        a = KO.cell(row=20+k,column=co).value; b = K.cell(row=33+k,column=c2).value
        a2= KO.cell(row=20+k,column=co+1).value; b2= K.cell(row=33+k,column=c2+1).value
        if a!=b or a2!=b2: miss.append(f'  множник {pl} #{k}: {a}/{a2} → {b}/{b2}')
    for k in range(4):
        a = KO.cell(row=42+k,column=co).value; b = K.cell(row=54+k,column=c2).value
        if a!=b: miss.append(f'  пакет {pl} #{k}: {a} → {b}')
    for k in range(2):
        a = KO.cell(row=51+k,column=co).value; b = K.cell(row=61+k,column=c2).value
        if a!=b: miss.append(f'  SPA {pl} #{k}: {a} → {b}')
print(f'\n4) Перенесення даних 14 прайс-листів: {"ВСЕ ЗБІГАЄТЬСЯ" if not miss else str(len(miss))+" розбіжностей"}')
for x in miss[:15]: print(x)
