from openpyxl.styles import Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from style import *
from names import *

HEAD = ['Категорія','Назва на Booking','Група','Опорна базова категорія','Мін. гостей',
        'Включено в базову','Макс. гостей','Роль',"SPA-візит: 1 од. / об'єкт"]
WID  = [40,38,20,34,12,15,13,14,26]

def build(wb, data):
    ws = wb.create_sheet('Довідник категорій')
    title(ws,'A1','ДОВІДНИК КАТЕГОРІЙ — РЕДАГУЙТЕ ЛИШЕ ПРИ ДОДАВАННІ / ЗМІНІ КАТЕГОРІЇ','A1:I1')
    banner(ws,'A2','У щоденній роботі цей лист не редагують. Він визначає місткість, опорну базову категорію та SPA-роль — саме з нього формули беруть правила. '
                   'Зміна «Макс. гостей» одразу прибирає або додає варіант розміщення в листах «Прайс», «Калькулятор» і «Контроль і порівняння».','A2:I2',34)
    for i,h in enumerate(HEAD):
        col_head(ws, f'{col_letter(i+1)}4', h)
        ws.column_dimensions[col_letter(i+1)].width = WID[i]
    ws.row_dimensions[4].height = 32

    for r,row in enumerate(data['cats'], start=5):
        band = (r % 2 == 1)
        for c,v in enumerate(row, start=1):
            cell = f'{col_letter(c)}{r}'
            if c in (5,6,7):                       # місткість — ввід
                inp(ws, cell, v, '0')
                ws[cell].alignment = Alignment(horizontal='center')
            else:
                out(ws, cell, v, band=band)
        ws[f'A{r}'].font = f(b=True)

    dv_role = DataValidation(type='list', formula1='"Базова,За матрицею"', allow_blank=True)
    dv_spa  = DataValidation(type='list', formula1='=SPA_РОЛЬ', allow_blank=True)
    dv_base = DataValidation(type='list', formula1='=БАЗА_КАТ', allow_blank=True)
    dv_num  = DataValidation(type='whole', operator='between', formula1='0', formula2='20',
                             allow_blank=True, showErrorMessage=True,
                             errorTitle='Некоректна місткість', error='Введіть ціле число від 0 до 20.')
    for dv,rng in ((dv_role,'H5:H22'),(dv_spa,'I5:I22'),(dv_base,'D5:D22'),(dv_num,'E5:G22')):
        ws.add_data_validation(dv); dv.add(rng)

    ws.freeze_panes = 'A5'
    ws.sheet_view.showGridLines = False
    return ws
