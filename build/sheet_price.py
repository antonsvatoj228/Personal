from openpyxl.styles import Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as gl
from style import *
from names import *

R_GRP, R_HEAD, R0 = 7, 8, 9              # групи колонок / шапка / перший рядок даних
R1 = R0 + N_CAT - 1                      # 26
WD0, WD1 = 4, 11                         # БУДНІ   D..K
SPACER   = 12                            # колонка-роздільник L
WE0, WE1 = 13, 20                        # ВИХІДНІ M..T
HELP0 = 21                               # службові колонки з U
GRP_END = (10, 15, 20)                   # останній рядок кожної групи категорій

PKG_DESC = ('=IF($B$5="RO","Лише проживання.",'
            'IF($B$5="BB","Проживання + сніданок за кожного гостя.",'
            'IF($B$5="BBSPA","Проживання + сніданок і SPA за кожного гостя.",'
            'IF($B$5="ROSPA","Проживання + 1 SPA-візит на обʼєкт.",'
            'IF($B$5="ROSPABB","Проживання + сніданки за кількістю гостей + 1 SPA-візит на обʼєкт.","")))))')

def build(wb, data):
    ws = wb.create_sheet('Прайс', 0)
    title(ws,'A1',None,'A1:S1')
    ws['A1'] = '="ПРАЙС BUKOVEL · "&$B$3&" · "&$U$11&" · Тариф: "&$B$4&" · Пакет: "&$B$5'
    ws.row_dimensions[2].height = 6

    for r,(lab,default) in enumerate(
            [('Активний прайс-лист', 'Осінь 2026 (00-10%)'),
             ('Тариф (знижка)',      'Rack Rate'),
             ('Пакет',               'ROSPA')], start=3):
        label(ws,f'A{r}',lab,bold=True)
        inp(ws,f'B{r}',default)
        ws.row_dimensions[r].height = 19

    ws['D3'] = '="Період дії: "&$U$11'
    ws['D4'] = ('="Знижка на проживання: "&TEXT($U$5,"0%")'
                '&IF($B$5="BBSPA"," — фіксована для пакета BBSPA; обраний тариф не застосовується","")')
    ws['D5'] = PKG_DESC
    for r in (3,4,5):
        ws[f'D{r}'].font = f(10, it=True, color='2A3B33')
        ws[f'D{r}'].fill = fill(GREEN_PL)
        ws[f'D{r}'].alignment = Alignment(vertical='center', indent=1)
        for c in range(5, WE1+1): ws.cell(row=r, column=c).fill = fill(GREEN_PL)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=WE1)

    banner(ws,'A6',None,f'A6:{gl(WE1)}6',30)
    ws['A6'] = ('=IF($U$3=0,"⚠ Прайс-лист не обрано або не знайдено — оберіть значення у клітинці B3.",'
                'IF($U$13=0,"⚠ У цьому прайс-листі ще не заповнені базові RO-ціни — сітка буде порожньою. '
                'Заповніть його блок на листі «Керування прайсом».",'
                '"Усі ціни — за ніч, за дорослих гостей, у гривні. Дитячі тарифи рахує лист «Калькулятор». '
                'З 3-го гостя додається фіксований крок; вихідні рахуються через множник вихідних."))')
    ws.conditional_formatting.add(f'A6:{gl(WE1)}6',
        FormulaRule(formula=['OR($U$3=0,$U$13=0)'], fill=fill(RED_SOFT), font=f(10,True,RED_INK), stopIfTrue=True))

    # шапка сітки
    for c0,c1,txt in ((WD0,WD1,'БУДНІ — кількість гостей'),(WE0,WE1,'ВИХІДНІ — кількість гостей')):
        block_head(ws,f'{gl(c0)}{R_GRP}',txt,f'{gl(c0)}{R_GRP}:{gl(c1)}{R_GRP}')
        ws[f'{gl(c0)}{R_GRP}'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[R_GRP].height = 20
    for c,h in ((1,'Категорія'),(2,'Назва на Booking'),(3,'Група')):
        col_head(ws,f'{gl(c)}{R_HEAD}',h)
    for c in range(WD0,WE1+1):
        if c == SPACER: continue
        n = (c - WD0 + 1) if c <= WD1 else (c - WE0 + 1)
        col_head(ws,f'{gl(c)}{R_HEAD}',n); ws[f'{gl(c)}{R_HEAD}'].number_format = '0'
    ws.row_dimensions[R_HEAD].height = 20

    for w,c in zip([46,40,17],'ABC'): ws.column_dimensions[c].width = w
    for c in range(WD0,WE1+1): ws.column_dimensions[gl(c)].width = 12.5
    ws.column_dimensions[gl(SPACER)].width = 2.6

    # глобальні службові клітинки
    g = {
      3:  '=IFERROR(MATCH($B$3,ПЛ_РЯДОК,0),0)',
      4:  '=IFERROR(SUMIFS(ТАРИФИ_ЗНИЖКИ,ТАРИФИ_НАЗВИ,$B$4),0)',
      5:  '=IF($B$5="BBSPA",ЗНИЖКА_BBSPA,$U$4)',
      6:  '=IF($U$3=0,0,IFERROR(INDEX(ПЕРЕМИКАЧІ,IF($B$5="BBSPA",2,1),$U$3),0))',
      7:  '=IF($U$3=0,0,IFERROR(INDEX(ПЕРЕМИКАЧІ,3,$U$3),0))',
      8:  '=IF(OR($B$5="BB",$B$5="BBSPA",$B$5="ROSPABB"),1,0)',
      9:  '=IF(OR($B$5="ROSPA",$B$5="ROSPABB"),1,0)',
      10: '=IF($U$8=0,"",IF($B$5="ROSPABB","BB",$B$5))',
      12: '=IFERROR(MATCH($B$3,ПРАЙС_ЛИСТИ,0),0)',
      11: ('=IF($U$12=0,"період не заданий",IFERROR(IF(OR(INDEX(ПЛ_ДАТА_З,$U$12)="",INDEX(ПЛ_ДАТА_ПО,$U$12)=""),'
           '"період не заданий",'
           'TEXT(DAY(INDEX(ПЛ_ДАТА_З,$U$12)),"00")&"."&TEXT(MONTH(INDEX(ПЛ_ДАТА_З,$U$12)),"00")&"."&YEAR(INDEX(ПЛ_ДАТА_З,$U$12))'
           '&" – "&'
           'TEXT(DAY(INDEX(ПЛ_ДАТА_ПО,$U$12)),"00")&"."&TEXT(MONTH(INDEX(ПЛ_ДАТА_ПО,$U$12)),"00")&"."&YEAR(INDEX(ПЛ_ДАТА_ПО,$U$12))),'
           '"період не заданий"))'),
      13: ('=IF($U$3=0,0,COUNT(INDEX(БАЗА_RO,1,$U$3),INDEX(БАЗА_RO,2,$U$3),'
           'INDEX(БАЗА_RO,3,$U$3),INDEX(БАЗА_RO,4,$U$3)))'),
    }
    HL = gl(HELP0)
    labels = {3:'колонка прайс-листа',4:'знижка тарифу',5:'знижка на проживання',6:'перемикач: пакет',
              7:'перемикач: SPA',8:'потрібен пакет',9:'потрібен SPA',10:'ключ пакета',
              11:'період (текст)',12:'рядок прайс-листа',13:'заповнених баз RO'}
    for r,fml in g.items():
        ws[f'{HL}{r}'] = fml
        ws[f'{HL}{r}'].font = f(9, color='6B7771')
        ws[f'{gl(HELP0+1)}{r}'] = labels[r]
        ws[f'{gl(HELP0+1)}{r}'].font = f(9, it=True, color='9AA6A0')
    ws[f'{HL}1'] = 'ТЕХНІЧНІ ДАНІ — НЕ РЕДАГУВАТИ'
    ws[f'{HL}1'].font = f(10, True, '6B7771')

    _rows(ws)
    _validation(ws)

    ws.freeze_panes = f'{gl(WD0)}{R0}'
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GREEN_D
    ws.print_area = f'A1:{gl(WE1)}{R1}'
    ws.print_title_rows = f'1:{R_HEAD}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.left.text = 'Phoenix Relax Park Bukovel'
    ws.oddFooter.center.text = '&A'
    ws.oddFooter.right.text = 'Сторінка &P з &N'
    return ws


def _rows(ws):
    cols = dict(zip(['ro','mult','wmult','maxg','incl','pkgd','pkge','spa','okd','oke'],
                    [gl(HELP0+2+i) for i in range(10)]))
    for i in range(N_CAT):
        r  = R0 + i
        dr = 5 + i                                   # рядок у «Довідник категорій»
        band = (i % 2 == 1)
        ws[f'A{r}'] = f"=IF('Довідник категорій'!A{dr}=\"\",\"\",'Довідник категорій'!A{dr})"
        ws[f'B{r}'] = f'=IF($A{r}="","",IFERROR(INDEX(ДОВ_BOOKING,MATCH($A{r},КАТЕГОРІЇ,0)),""))'
        ws[f'C{r}'] = f'=IF($A{r}="","",IFERROR(INDEX(ДОВ_ГРУПА,MATCH($A{r},КАТЕГОРІЇ,0)),""))'
        out(ws,f'A{r}',None,band=band); ws[f'A{r}'].font = f(b=True)
        out(ws,f'B{r}',None,band=band); out(ws,f'C{r}',None,band=band)
        ws[f'C{r}'].alignment = Alignment(horizontal='center')

        m = cols
        ws[f'{m["ro"]}{r}']    = (f'=IF(OR($U$3=0,$A{r}=""),0,IFERROR(INDEX(БАЗА_RO,'
                                 f'MATCH(INDEX(ДОВ_ОПОРА,MATCH($A{r},КАТЕГОРІЇ,0)),БАЗА_КАТ,0),$U$3),0))')
        ws[f'{m["mult"]}{r}']  = f'=IF(OR($U$3=0,$A{r}=""),0,IFERROR(INDEX(МНОЖНИКИ,MATCH($A{r},МНОЖ_КАТ,0),$U$3),0))'
        ws[f'{m["wmult"]}{r}'] = f'=IF(OR($U$3=0,$A{r}=""),0,IFERROR(INDEX(МНОЖНИКИ,MATCH($A{r},МНОЖ_КАТ,0),$U$3+1),0))'
        ws[f'{m["maxg"]}{r}']  = f'=IF($A{r}="",0,IFERROR(INDEX(ДОВ_МАКС,MATCH($A{r},КАТЕГОРІЇ,0)),0))'
        ws[f'{m["incl"]}{r}']  = f'=IF($A{r}="",0,IFERROR(INDEX(ДОВ_ВКЛ,MATCH($A{r},КАТЕГОРІЇ,0)),0))'
        ws[f'{m["pkgd"]}{r}']  = (f'=IF(OR($U$3=0,$U$8=0),0,IFERROR(INDEX(ПАКЕТИ_ЦІНИ,'
                                 f'MATCH($U$10&"|Будні",ПАКЕТИ_КЛЮЧ,0),$U$3),0))')
        ws[f'{m["pkge"]}{r}']  = (f'=IF(OR($U$3=0,$U$8=0),0,IFERROR(INDEX(ПАКЕТИ_ЦІНИ,'
                                 f'MATCH($U$10&"|Вихідні",ПАКЕТИ_КЛЮЧ,0),$U$3),0))')
        ws[f'{m["spa"]}{r}']   = (f'=IF(OR($U$3=0,$U$9=0,$A{r}=""),0,IFERROR(INDEX(SPA_ЦІНИ,'
                                 f'MATCH(INDEX(ДОВ_SPA,MATCH($A{r},КАТЕГОРІЇ,0)),SPA_РОЛЬ,0),$U$3),0))')
        ws[f'{m["okd"]}{r}']   = (f'=IF(OR($A{r}="",{m["ro"]}{r}=0,{m["mult"]}{r}=0,'
                                 f'AND($U$8=1,{m["pkgd"]}{r}=0),AND($U$9=1,{m["spa"]}{r}=0)),0,1)')
        ws[f'{m["oke"]}{r}']   = (f'=IF(OR({m["okd"]}{r}=0,{m["wmult"]}{r}=0,'
                                 f'AND($U$8=1,{m["pkge"]}{r}=0)),0,1)')
        for key in m.values():
            ws[f'{key}{r}'].font = f(9, color='6B7771')

        for c in range(WD0, WE1+1):
            if c == SPACER: continue
            wknd = c >= WE0
            nref = f'{gl(c)}${R_HEAD}'
            okc  = m['oke'] if wknd else m['okd']
            pkg  = m['pkge'] if wknd else m['pkgd']
            core = (f'{m["ro"]}{r}*{m["mult"]}{r}' + (f'*{m["wmult"]}{r}' if wknd else ''))
            ws.cell(row=r, column=c).value = (
                f'=IF(OR(${okc}{r}=0,{nref}>${m["maxg"]}{r}),"",'
                f'ROUND((({core}+MAX(0,{nref}-${m["incl"]}{r})*КРОК_ГОСТЯ)*(1-$U$5)'
                f'+{nref}*${pkg}{r}*(1-$U$5*$U$6)'
                f'+${m["spa"]}{r}*(1-$U$5*$U$7))/ОКРУГЛЕННЯ,0)*ОКРУГЛЕННЯ)')
            out(ws, f'{gl(c)}{r}', None, MONEY, band=band)
            ws.cell(row=r, column=c).alignment = Alignment(horizontal='right')

    # ── межі, що допомагають читати сітку ──
    for r in range(R0, R1 + 1):
        edge(ws, f'C{r}', right=True)                     # опис | ціни
        edge(ws, f'{gl(WD1)}{r}', right=True)             # кінець буднів
        edge(ws, f'{gl(WE0)}{r}', left=True)              # початок вихідних
        if r in GRP_END:                                  # межа між групами категорій
            for c in range(1, WE1 + 1):
                if c == SPACER: continue
                edge(ws, f'{gl(c)}{r}', bottom=True)
    edge(ws, f'C{R_HEAD}', right=True)
    edge(ws, f'{gl(WD1)}{R_HEAD}', right=True)
    edge(ws, f'{gl(WE0)}{R_HEAD}', left=True)

    # базові категорії — виділення
    ws.conditional_formatting.add(f'A{R0}:C{R1}',
        FormulaRule(formula=[f'IFERROR(INDEX(ДОВ_РОЛЬ,MATCH($A{R0},КАТЕГОРІЇ,0))="Базова",FALSE)'],
                    font=f(11, True, RED_INK), stopIfTrue=False))
    # службові колонки — згорнути й приховати
    for c in range(HELP0, HELP0 + 13):
        ws.column_dimensions[gl(c)].hidden = True
        ws.column_dimensions[gl(c)].outlineLevel = 1


def _validation(ws):
    for rng, src in (('B3','=ПРАЙС_ЛИСТИ'), ('B4','=ТАРИФИ_НАЗВИ'), ('B5','=СПИСОК_ПАКЕТІВ')):
        dv = DataValidation(type='list', formula1=src, allow_blank=False,
                            showErrorMessage=True, errorTitle='Оберіть зі списку',
                            error='Значення має бути з випадаючого списку — інакше прайс не порахується.')
        ws.add_data_validation(dv); dv.add(rng)
