from openpyxl.styles import Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.utils import get_column_letter as gl
from style import *
from names import *

# ═══════════════════ МОДЕЛЮВАННЯ ═══════════════════
def build_model(wb, data):
    ws = wb.create_sheet('Моделювання')
    R0 = 8; R1 = R0 + N_CAT - 1
    title(ws,'A1','МОДЕЛЮВАННЯ ЦІН — ДВА СЦЕНАРІЇ У ПОРІВНЯННІ','A1:U1')
    label(ws,'A3','Прайс-лист-основа',bold=True); inp(ws,'B3','Осінь 2026 (00-10%)')
    dv = DataValidation(type='list', formula1='=ПРАЙС_ЛИСТИ', allow_blank=False)
    ws.add_data_validation(dv); dv.add('B3')
    banner(ws,'A4','Жовті колонки — ввід множників. Сценарії нічого не змінюють у діючому прайсі: щоб застосувати сценарій, '
                   'скопіюйте його множники у блок відповідного прайс-листа на листі «Керування прайсом». '
                   'Ціни рахуються для RO 1–2 особи — це чистий важіль сходинки між категоріями.','A4:U4',34)

    GROUPS = [(4,7,'ПОТОЧНО'),(8,13,'СЦЕНАРІЙ 1'),(14,19,'СЦЕНАРІЙ 2'),(20,21,'СЦЕНАРІЙ 2 vs 1')]
    for c0,c1,t in GROUPS:
        block_head(ws,f'{gl(c0)}6',t,f'{gl(c0)}6:{gl(c1)}6')
        ws[f'{gl(c0)}6'].alignment = Alignment(horizontal='center',vertical='center')
    HEAD = ['Категорія','Опорна базова','RO база, грн',
            'Множник','Множник вих.','Будні, грн','Вихідні, грн',
            'Множник','Множник вих.','Будні, грн','Вихідні, грн','Δ грн буд.','Δ % буд.',
            'Множник','Множник вих.','Будні, грн','Вихідні, грн','Δ грн буд.','Δ % буд.',
            'Δ грн буд.','Δ % буд.']
    for i,h in enumerate(HEAD): col_head(ws,f'{gl(i+1)}7',h)
    ws.row_dimensions[7].height = 32
    for w,c in zip([38,30,13],'ABC'): ws.column_dimensions[c].width = w
    for c in range(4,22): ws.column_dimensions[gl(c)].width = 12

    ws['W2'] = '=IFERROR(MATCH($B$3,ПЛ_РЯДОК,0),0)'; ws['W2'].font = f(9,color='6B7771')
    ws['X2'] = 'колонка прайс-листа'; ws['X2'].font = f(9,it=True,color='9AA6A0')
    ws['W1'] = 'ТЕХНІЧНІ ДАНІ'; ws['W1'].font = f(10,True,'6B7771')
    for c in ('W','X'): ws.column_dimensions[c].hidden = True

    for i in range(N_CAT):
        r, dr, band = R0+i, 5+i, i % 2 == 1
        src = data['pricelists'][2]['mult'][i] if len(data['pricelists']) > 2 else (None,None)
        m = f'MATCH($A{r},КАТЕГОРІЇ,0)'
        ws[f'A{r}'] = f"=IF('Довідник категорій'!A{dr}=\"\",\"\",'Довідник категорій'!A{dr})"
        ws[f'B{r}'] = f'=IF($A{r}="","",IFERROR(INDEX(ДОВ_ОПОРА,{m}),""))'
        ws[f'C{r}'] = (f'=IF(OR($W$2=0,$A{r}=""),"",IFERROR(INDEX(БАЗА_RO,'
                       f'MATCH($B{r},БАЗА_КАТ,0),$W$2),""))')
        ws[f'D{r}'] = f'=IF(OR($W$2=0,$A{r}=""),"",IFERROR(INDEX(МНОЖНИКИ,MATCH($A{r},МНОЖ_КАТ,0),$W$2),""))'
        ws[f'E{r}'] = f'=IF(OR($W$2=0,$A{r}=""),"",IFERROR(INDEX(МНОЖНИКИ,MATCH($A{r},МНОЖ_КАТ,0),$W$2+1),""))'
        ws[f'F{r}'] = f'=IF(OR($C{r}="",$D{r}=""),"",ROUND($C{r}*$D{r}/ОКРУГЛЕННЯ,0)*ОКРУГЛЕННЯ)'
        ws[f'G{r}'] = f'=IF(OR($C{r}="",$D{r}="",$E{r}=""),"",ROUND($C{r}*$D{r}*$E{r}/ОКРУГЛЕННЯ,0)*ОКРУГЛЕННЯ)'
        out(ws,f'A{r}',None,band=band); ws[f'A{r}'].font = f(b=True)
        out(ws,f'B{r}',None,band=band); ws[f'B{r}'].font = f(10)
        out(ws,f'C{r}',None,MONEY,band=band)
        out(ws,f'D{r}',None,MULT,band=band); out(ws,f'E{r}',None,MULT,band=band)
        out(ws,f'F{r}',None,MONEY,band=band); out(ws,f'G{r}',None,MONEY,band=band)

        for base,alt in ((8,False),(14,True)):
            cm, cw, cb, cv, cd, cp = (gl(base+k) for k in range(6))
            inp(ws,f'{cm}{r}', src[0], MULT, alt=alt); inp(ws,f'{cw}{r}', src[1], MULT, alt=alt)
            ws[f'{cb}{r}'] = f'=IF(OR($C{r}="",{cm}{r}=""),"",ROUND($C{r}*{cm}{r}/ОКРУГЛЕННЯ,0)*ОКРУГЛЕННЯ)'
            ws[f'{cv}{r}'] = f'=IF(OR($C{r}="",{cm}{r}="",{cw}{r}=""),"",ROUND($C{r}*{cm}{r}*{cw}{r}/ОКРУГЛЕННЯ,0)*ОКРУГЛЕННЯ)'
            ws[f'{cd}{r}'] = f'=IF(OR({cb}{r}="",$F{r}=""),"",{cb}{r}-$F{r})'
            ws[f'{cp}{r}'] = f'=IF(OR({cb}{r}="",$F{r}="",$F{r}=0),"",{cb}{r}/$F{r}-1)'
            out(ws,f'{cb}{r}',None,MONEY,band=band); out(ws,f'{cv}{r}',None,MONEY,band=band)
            out(ws,f'{cd}{r}',None,MONEY,band=band); out(ws,f'{cp}{r}',None,PCT,band=band)
        ws[f'T{r}'] = f'=IF(OR(P{r}="",J{r}=""),"",P{r}-J{r})'
        ws[f'U{r}'] = f'=IF(OR(P{r}="",J{r}="",J{r}=0),"",P{r}/J{r}-1)'
        out(ws,f'T{r}',None,MONEY,band=band); out(ws,f'U{r}',None,PCT,band=band)

    rs = R1 + 2
    label(ws,f'A{rs}','ПІДСУМОК — середня зміна до поточного прайсу',bold=True)
    ws[f'M{rs}'] = f'=IFERROR(AVERAGE(M{R0}:M{R1}),"")'
    ws[f'S{rs}'] = f'=IFERROR(AVERAGE(S{R0}:S{R1}),"")'
    ws[f'U{rs}'] = f'=IFERROR(AVERAGE(U{R0}:U{R1}),"")'
    for c in ('M','S','U'):
        out(ws,f'{c}{rs}',None,PCT,bold=True); ws[f'{c}{rs}'].fill = fill(GREEN_PL)
    for rng in (f'L{R0}:L{R1}',f'M{R0}:M{R1}',f'R{R0}:R{R1}',f'S{R0}:S{R1}',f'T{R0}:T{R1}',f'U{R0}:U{R1}'):
        ws.conditional_formatting.add(rng, CellIsRule(operator='lessThan', formula=['0'], font=f(11,color=RED_INK)))
    ws.conditional_formatting.add(f'A{R0}:C{R1}',
        FormulaRule(formula=[f'IFERROR(INDEX(ДОВ_РОЛЬ,MATCH($A{R0},КАТЕГОРІЇ,0))="Базова",FALSE)'],
                    font=f(11,True,RED_INK), stopIfTrue=False))
    dvp = DataValidation(type='decimal', operator='between', formula1='0', formula2='10', allow_blank=True,
                         showErrorMessage=True, errorTitle='Множник',
                         error='Множник має бути числом від 0 до 10.')
    ws.add_data_validation(dvp); dvp.add(f'H{R0}:I{R1}'); dvp.add(f'N{R0}:O{R1}')
    ws.freeze_panes = f'D{R0}'
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = '8A7128'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


# ═══════════════════ МАТРИЦЯ КАТЕГОРІЙ ═══════════════════
def build_matrix(wb, data):
    ws = wb.create_sheet('Матриця категорій')
    R0 = 7; R1 = R0 + N_CAT - 1
    title(ws,'A1','МАТРИЦЯ КАТЕГОРІЙ — ПЕРЕГЛЯД АКТИВНОГО ПРАЙСУ','A1:H1')
    label(ws,'A3','Активний прайс-лист',bold=True)
    ws['B3'] = "='Прайс'!$B$3"; ws['B3'].fill = fill(GREEN_PL); ws['B3'].font = f(b=True,color=GREEN_D)
    banner(ws,'A4','Лист лише показує результат — редагування множників і цін виконується на «Керування прайсом», '
                   'моделювання — на листі «Моделювання». Вихідні 1–2 ос. = RO будні × множник вихідних; '
                   'з 3-го гостя у прайсі додається фіксований крок без множника.','A4:H4',34)
    for i,h in enumerate(['Категорія','Опорна базова категорія','Множник ціни','Множник вихідних',
                          'RO будні 1–2, грн','RO вихідні 1–2, грн','Δ грн до попередньої','Δ грн буд.→вих.']):
        col_head(ws,f'{gl(i+1)}6',h)
    ws.row_dimensions[6].height = 32
    for w,c in zip([40,32,13,15,15,17,16,15],'ABCDEFGH'): ws.column_dimensions[c].width = w
    ws['J2'] = "=IFERROR(MATCH('Прайс'!$B$3,ПЛ_РЯДОК,0),0)"; ws['J2'].font = f(9,color='6B7771')
    ws['J1'] = 'ТЕХНІЧНІ ДАНІ'; ws['J1'].font = f(10,True,'6B7771')
    ws.column_dimensions['J'].hidden = True

    for i in range(N_CAT):
        r, dr, band = R0+i, 5+i, i % 2 == 1
        m = f'MATCH($A{r},КАТЕГОРІЇ,0)'
        ws[f'A{r}'] = f"=IF('Довідник категорій'!A{dr}=\"\",\"\",'Довідник категорій'!A{dr})"
        ws[f'B{r}'] = f'=IF($A{r}="","",IFERROR(INDEX(ДОВ_ОПОРА,{m}),""))'
        ws[f'C{r}'] = f'=IF(OR($J$2=0,$A{r}=""),"",IFERROR(INDEX(МНОЖНИКИ,MATCH($A{r},МНОЖ_КАТ,0),$J$2),""))'
        ws[f'D{r}'] = f'=IF(OR($J$2=0,$A{r}=""),"",IFERROR(INDEX(МНОЖНИКИ,MATCH($A{r},МНОЖ_КАТ,0),$J$2+1),""))'
        ws[f'E{r}'] = (f'=IF(OR($J$2=0,$A{r}="",$C{r}=""),"",IFERROR(ROUND(INDEX(БАЗА_RO,'
                       f'MATCH($B{r},БАЗА_КАТ,0),$J$2)*$C{r}/ОКРУГЛЕННЯ,0)*ОКРУГЛЕННЯ,""))')
        ws[f'F{r}'] = (f'=IF(OR($J$2=0,$A{r}="",$C{r}="",$D{r}=""),"",IFERROR(ROUND(INDEX(БАЗА_RO,'
                       f'MATCH($B{r},БАЗА_КАТ,0),$J$2)*$C{r}*$D{r}/ОКРУГЛЕННЯ,0)*ОКРУГЛЕННЯ,""))')
        prev = r - 1
        ws[f'G{r}'] = ('' if i == 0 else
                       f'=IF(OR($E{r}="",$E{prev}="",$B{r}<>$B{prev}),"",$E{r}-$E{prev})')
        ws[f'H{r}'] = f'=IF(OR($E{r}="",$F{r}=""),"",$F{r}-$E{r})'
        out(ws,f'A{r}',None,band=band); ws[f'A{r}'].font = f(b=True)
        out(ws,f'B{r}',None,band=band); ws[f'B{r}'].font = f(10)
        out(ws,f'C{r}',None,MULT,band=band); out(ws,f'D{r}',None,MULT,band=band)
        out(ws,f'E{r}',None,MONEY,band=band); out(ws,f'F{r}',None,MONEY,band=band)
        out(ws,f'G{r}',None,MONEY,band=band); out(ws,f'H{r}',None,MONEY,band=band)
    ws.conditional_formatting.add(f'A{R0}:B{R1}',
        FormulaRule(formula=[f'IFERROR(INDEX(ДОВ_РОЛЬ,MATCH($A{R0},КАТЕГОРІЇ,0))="Базова",FALSE)'],
                    fill=fill(RED_SOFT), font=f(11,True,RED_INK), stopIfTrue=False))
    ws.freeze_panes = f'C{R0}'
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = 'landscape'
    return ws


# ═══════════════════ БАЗОВІ ЦІНИ ═══════════════════
def build_base(wb, data):
    ws = wb.create_sheet('Базові ціни')
    title(ws,'A1','БАЗОВІ ЦІНИ — ПЕРЕГЛЯД АКТИВНОГО ПРАЙСУ','A1:E1')
    label(ws,'A3','Активний прайс-лист',bold=True)
    ws['B3'] = "='Прайс'!$B$3"; ws['B3'].fill = fill(GREEN_PL); ws['B3'].font = f(b=True,color=GREEN_D)
    banner(ws,'A4','Чотири опорні категорії, з яких множниками формуються всі інші. Редагування — на «Керування прайсом».','A4:E4',26)
    for i,h in enumerate(['Базова категорія','RO 1–2, грн','Множник вихідних','Вихідні 1–2, грн','Статус']):
        col_head(ws,f'{gl(i+1)}6',h)
    for w,c in zip([34,16,18,17,14],'ABCDE'): ws.column_dimensions[c].width = w
    ws['G2'] = "=IFERROR(MATCH('Прайс'!$B$3,ПЛ_РЯДОК,0),0)"; ws['G2'].font = f(9,color='6B7771')
    ws.column_dimensions['G'].hidden = True
    for i,name in enumerate(data['basecats']):
        r = 7 + i
        ws[f'A{r}'] = name
        ws[f'B{r}'] = f'=IF($G$2=0,"",IFERROR(INDEX(БАЗА_RO,MATCH($A{r},БАЗА_КАТ,0),$G$2),""))'
        ws[f'C{r}'] = f'=IF($G$2=0,"",IFERROR(INDEX(МНОЖНИКИ,MATCH($A{r},МНОЖ_КАТ,0),$G$2+1),""))'
        ws[f'D{r}'] = f'=IF(OR($B{r}="",$C{r}=""),"",ROUND($B{r}*$C{r}/ОКРУГЛЕННЯ,0)*ОКРУГЛЕННЯ)'
        ws[f'E{r}'] = f'=IF($B{r}="","ПЕРЕВІРИТИ","OK")'
        out(ws,f'A{r}',None,bold=True); out(ws,f'B{r}',None,MONEY); out(ws,f'C{r}',None,MULT)
        out(ws,f'D{r}',None,MONEY);    out(ws,f'E{r}')
        ws[f'E{r}'].alignment = Alignment(horizontal='center')
    ws.conditional_formatting.add('E7:E10', FormulaRule(formula=['$E7="ПЕРЕВІРИТИ"'],
                                  fill=fill(RED_SOFT), font=f(11,True,RED_INK)))
    ws.sheet_view.showGridLines = False
    return ws


# ═══════════════════ ЖУРНАЛ ЗМІН ═══════════════════
def build_log(wb, data, version, today):
    ws = wb.create_sheet('Журнал змін')
    title(ws,'A1','ЖУРНАЛ ЗМІН','A1:F1')
    banner(ws,'A2','Заповнюється вручну. Правило: діючий прайс-лист не перезаписуємо — переводимо його у статус «Архів» '
                   'і створюємо новий рядок на «Керування прайсом». Тоді порівняння рік/рік лишається можливим.','A2:F2',30)
    for i,h in enumerate(['Дата','Автор','Лист','Що змінено','Було','Стало']):
        col_head(ws,f'{gl(i+1)}4',h)
    for w,c in zip([13,20,26,46,20,20],'ABCDEF'): ws.column_dimensions[c].width = w
    rows = [
      (today,'Claude (міграція v2)','Довідник категорій',
       'Апартаменти Люкс 25 м²: макс. гостей','3','2'),
      (today,'Claude (міграція v2)','Керування прайсом',
       'Перемикачі знижки на пакети та SPA (за замовчуванням 0 — поведінка як раніше)','—','0 / 0 / 0'),
      (today,'Claude (міграція v2)','Прайс',
       'Будні й вихідні показуються одночасно; тип дня як перемикач прибрано','перемикач','дві групи колонок'),
    ]
    for i,row in enumerate(rows):
        r = 5 + i
        for j,v in enumerate(row):
            out(ws,f'{gl(j+1)}{r}',v,DATE if j==0 else None,band=(i%2==1))
            ws[f'{gl(j+1)}{r}'].alignment = Alignment(wrap_text=True, vertical='top')
    for r in range(8, 60):
        for j in range(6):
            inp(ws,f'{gl(j+1)}{r}',None, DATE if j==0 else None)
    ws.freeze_panes = 'A5'
    ws.sheet_view.showGridLines = False
    return ws


# ═══════════════════ ДОДАТКОВІ ПОСЛУГИ ═══════════════════
def build_extra(wb, data):
    ws = wb.create_sheet('Додаткові послуги')
    title(ws,'A1','ДОДАТКОВІ ПОСЛУГИ — ІНФОРМАЦІЙНО','A1:E1')
    banner(ws,'A2','Цей лист НЕ бере участі в розрахунку прайсу. Ціни сніданків і SPA, які реально рахує система, '
                   'задаються на листі «Керування прайсом» у блоках «ПАКЕТИ BB / BBSPA» та «ФІКСОВАНИЙ SPA-ВІЗИТ». '
                   'Тут — довідкові умови продажу та взаєморозрахунків.','A2:E2',40)
    for w,c in zip([30,22,22,16,26],'ABCDE'): ws.column_dimensions[c].width = w
    r = 4
    for row in data['dop']:
        if all(v in (None,'') for v in row): r += 1; continue
        for j,v in enumerate(row):
            if v is None: continue
            cell = f'{gl(j+1)}{r}'
            ws[cell] = v
            ws[cell].font = f()
            if isinstance(v,(int,float)) and j > 0:
                ws[cell].number_format = MONEY
        first = str(row[0] or '')
        if first.startswith('Додаткові послуги') or first in ('Буковель','BBSPA'):
            block_head(ws,f'A{r}',first,f'A{r}:E{r}')
        elif first == 'Назва послуги ':
            for j in range(5):
                if row[j] is not None: col_head(ws,f'{gl(j+1)}{r}',row[j])
        r += 1
    ws.sheet_view.showGridLines = False
    return ws
