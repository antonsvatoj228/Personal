from openpyxl.styles import Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from style import *
from names import *

PKG_LIST  = ['RO','BB','BBSPA','ROSPA','ROSPABB']
DAY_LIST  = ['Будні','Вихідні']
STAT_LIST = ['Активний','Чернетка','Архів']
SW_LABELS = ['Знижка на пакет BB (сніданки)',
             'Знижка на пакет BBSPA',
             'Знижка на фіксований SPA-візит (ROSPA / ROSPABB)']

def build(wb, data):
    ws = wb.create_sheet('Керування прайсом')
    title(ws,'A1','КЕРУВАННЯ ПРАЙСОМ — PHOENIX RELAX PARK BUKOVEL','A1:G1')
    banner(ws,'A2','Три кроки: 1) додайте прайс-лист у наступному вільному рядку зліва; 2) заповніть його блок праворуч; 3) перевірте результат у листі «Прайс». '
                   'Жовті клітинки — ввід, решта рахується автоматично.','A2:G2',34)

    # ── 1. Таблиця прайс-листів ────────────────────────────────────────────
    label(ws,'A4','1. ПРАЙС-ЛИСТИ / ПЕРІОДИ',bold=True)
    ws['A4'].font = f(11,True,GREEN_D)
    for c,h in zip('ABCDE',['Прайс-лист','Дата з','Дата по','Статус','Примітка / оновлено']):
        col_head(ws,f'{c}5',h)
    ws.row_dimensions[5].height = 30
    for w,c in zip([30,12,12,12,26],'ABCDE'):
        ws.column_dimensions[c].width = w

    for i in range(N_PL):
        r = R_PL0 + i
        src = data['pricelists'][i] if i < len(data['pricelists']) else None
        inp(ws,f'A{r}', src['name'].strip() if src else None)
        inp(ws,f'B{r}', src.get('date_from') if src else None, DATE)
        inp(ws,f'C{r}', src.get('date_to')   if src else None, DATE)
        inp(ws,f'D{r}', src['status'] if src else None)
        inp(ws,f'E{r}', None)
        ws[f'D{r}'].alignment = Alignment(horizontal='center')

    dv_stat = DataValidation(type='list', formula1='=СПИСОК_СТАТУСІВ', allow_blank=True)
    ws.add_data_validation(dv_stat); dv_stat.add(f'D{R_PL0}:D{R_PL_LAST}')
    dv_date = DataValidation(type='date', operator='between',
                             formula1='DATE(2020,1,1)', formula2='DATE(2040,12,31)',
                             allow_blank=True, showErrorMessage=True,
                             errorTitle='Некоректна дата', error='Введіть справжню дату у діапазоні 2020–2040.')
    ws.add_data_validation(dv_date); dv_date.add(f'B{R_PL0}:C{R_PL_LAST}')

    # ── 2. Глобальні параметри ─────────────────────────────────────────────
    label(ws,'F4','2. ГЛОБАЛЬНІ ПАРАМЕТРИ',bold=True); ws['F4'].font = f(11,True,GREEN_D)
    ws.column_dimensions['F'].width = 46
    ws.column_dimensions['G'].width = 24
    glob = [
        ('F5','Активний прайс-лист (показується у «Прайс»)', "='Прайс'!$B$3",      None,  False),
        ('F6','Крок за 3-го й кожного наступного гостя, грн', data['globals']['G6'], MONEY, True),
        ('F7','Крок округлення, грн',                          data['globals']['G7'], MONEY, True),
        ('F8','Знижка на проживання для пакета BBSPA',         data['globals']['G8'], '0%',  True),
        ('F9','Проживання з тваринами, грн / ніч',             data['globals']['G9'], MONEY, True),
        ('F10','Еталонний прайс-лист для Δ (варіант B)',        data['pricelists'][1]['name'].strip(), None, True),
        ('F11','Вихідні починаються з дня тижня (1=Пн … 7=Нд)', 5,                   '0',   True),
    ]
    for addr,lab,val,fmt,is_input in glob:
        label(ws,addr,lab,wrap=True)
        g = 'G'+addr[1:]
        if is_input: inp(ws,g,val,fmt)
        else:        out(ws,g,val,fmt,bold=True)
    ws['G5'].fill = fill(GREEN_PL); ws['G5'].font = f(b=True,color=GREEN_D)
    dv_et = DataValidation(type='list', formula1='=ПРАЙС_ЛИСТИ', allow_blank=True)
    ws.add_data_validation(dv_et); dv_et.add('G10')
    dv_r  = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0.01',
                           allow_blank=False, showErrorMessage=True, errorTitle='Крок округлення',
                           error='Крок округлення має бути не менший за 0,01 грн.')
    ws.add_data_validation(dv_r); dv_r.add('G7')
    dv_wd = DataValidation(type='whole', operator='between', formula1='1', formula2='7',
                           allow_blank=False, showErrorMessage=True, errorTitle='День тижня',
                           error='Введіть число від 1 (понеділок) до 7 (неділя).')
    ws.add_data_validation(dv_wd); dv_wd.add('G11')

    # ── 3. Тарифи ──────────────────────────────────────────────────────────
    label(ws,'F12','3. ТАРИФИ — ЗНИЖКИ НА ПРОЖИВАННЯ',bold=True); ws['F12'].font = f(11,True,GREEN_D)
    col_head(ws,f'F{R_TAR0-1}','Тариф'); col_head(ws,f'G{R_TAR0-1}','Знижка')
    for i in range(N_TAR):
        r = R_TAR0 + i
        src = data['tariffs'][i] if i < len(data['tariffs']) else (None,None)
        inp(ws,f'F{r}', src[0]); inp(ws,f'G{r}', src[1], '0%')

    # ── 4. Довідники для випадаючих списків ────────────────────────────────
    label(ws,'F23','4. ДОВІДНИКИ ДЛЯ ВИПАДАЮЧИХ СПИСКІВ',bold=True); ws['F23'].font = f(11,True,GREEN_D)
    col_head(ws,'F24','Пакети')
    for i,v in enumerate(PKG_LIST):  out(ws,f'F{25+i}',v)
    col_head(ws,'F31','Типи дня')
    for i,v in enumerate(DAY_LIST):  out(ws,f'F{32+i}',v)
    col_head(ws,'F35','Статуси')
    for i,v in enumerate(STAT_LIST): out(ws,f'F{36+i}',v)

    # ── Права частина: підписи рядків блоків ───────────────────────────────
    ws.column_dimensions['H'].width = 44
    lbl = lambda r,t,bold=True: (label(ws,f'H{r}',t,bold=bold,wrap=True))
    block_head(ws,f'H{R_BLK_NAME}','ПРАЙС-ЛИСТ')
    lbl(R_BLK_PERIOD,'Період дії',False); lbl(R_BLK_STATUS,'Статус',False); lbl(R_BLK_UPD,'Оновлено',False)
    block_head(ws,'H10','БАЗОВІ RO-ЦІНИ — 1–2 ОСОБИ')
    col_head(ws,'H11','Базова категорія')
    for i,n in enumerate(data['basecats']): out(ws,f'H{R_BASE0+i}',n,bold=True)

    block_head(ws,'H17','Δ МІЖ ОПОРНИМИ КАТЕГОРІЯМИ  (варіант A)')
    col_head(ws,'H18','Порівняння')
    DA_PAIRS = [(1,0),(2,0),(3,2)]           # (кого, з ким) за індексами basecats
    for i,(a,b) in enumerate(DA_PAIRS):
        out(ws,f'H{R_DA0+i}', f"{data['basecats'][a]}  ↔  {data['basecats'][b]}")
        ws[f'H{R_DA0+i}'].font = f(10)

    block_head(ws,'H23','Δ ДО ЕТАЛОННОГО ПРАЙС-ЛИСТА  (варіант B)')
    ws['H24'] = '=IF(ЕТАЛОН_ПЛ="","еталон не обрано — задайте G10","Еталон: "&ЕТАЛОН_ПЛ)'
    ws['H24'].font = f(10, it=True, color=GREEN_D); ws['H24'].fill = fill(GREEN_PL)
    col_head(ws,'H25','Базова категорія')
    for i,n in enumerate(data['basecats']): out(ws,f'H{R_DB0+i}',n,bold=True)

    block_head(ws,'H31','МАТРИЦЯ КАТЕГОРІЙ — МНОЖНИКИ')
    col_head(ws,'H32','Категорія')
    for i,n in enumerate(data['matcats']): out(ws,f'H{R_MULT0+i}',n)

    block_head(ws,'H52','ПАКЕТИ BB / BBSPA — ЗА 1 ГОСТЯ')
    col_head(ws,'H53','Пакет / тип дня')
    for i,n in enumerate(data['pkgkeys']): out(ws,f'H{R_PKG0+i}',n,bold=True)

    block_head(ws,'H59',"ФІКСОВАНИЙ SPA-ВІЗИТ ДЛЯ ROSPA / ROSPABB — 1 ОД. НА ОБ'ЄКТ")
    col_head(ws,'H60','SPA-рівень / категорії')
    for i,n in enumerate(data['sparoles']): out(ws,f'H{R_SPA0+i}',n,bold=True)

    block_head(ws,'H64','ПЕРЕМИКАЧІ ЗНИЖКИ   (1 = знижка тарифу діє, 0 = не діє)')
    col_head(ws,'H65','На що діє знижка тарифу')
    for i,n in enumerate(SW_LABELS): out(ws,f'H{R_SW0+i}',n)

    label(ws,'H70','Логіка: RO-ціна задається лише для 4 опорних категорій, решта формуються множниками. '
                   'З 3-го гостя додається фіксований крок (G6). BB / BBSPA додають пакет за кожного гостя, '
                   'ROSPA — один SPA-візит на об\'єкт, ROSPABB — сніданки за кількістю гостей + один SPA-візит.',wrap=True)
    ws['H70'].font = f(10, it=True)
    ws.merge_cells('H70:N70'); ws.row_dimensions[70].height = 44

    # ── Блоки прайс-листів ─────────────────────────────────────────────────
    _blocks(ws, data)

    ws.freeze_panes = 'I5'
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GREEN_D
    ws.page_setup.orientation = 'landscape'
    return ws


def _blocks(ws, data):
    for i in range(N_PL):
        c  = pl_col(i);  L, L2 = col_letter(c), col_letter(c+1)
        r  = R_PL0 + i
        src = data['pricelists'][i] if i < len(data['pricelists']) else None
        ws.column_dimensions[L].width  = 15
        ws.column_dimensions[L2].width = 17

        # технічний рядок 2 — саме його читає MATCH
        for cc in (L, L2):
            ws[f'{cc}{R_TECH}'] = f'=IF($A${r}="","",$A${r})'
            ws[f'{cc}{R_TECH}'].font = f(8, color='9AA6A0')
            ws[f'{cc}{R_TECH}'].fill = fill(GREY_T)

        # шапка блоку
        ws[f'{L}{R_BLK_NAME}'] = f'=IF($A${r}="","",$A${r})'
        ws[f'{L}{R_BLK_NAME}'].font = f(11, True, 'FFFFFF')
        ws[f'{L}{R_BLK_NAME}'].fill = fill(GREEN_D)
        ws[f'{L}{R_BLK_NAME}'].alignment = Alignment(vertical='center', wrap_text=True, indent=1)
        ws[f'{L2}{R_BLK_NAME}'].fill = fill(GREEN_D)
        ws.merge_cells(f'{L}{R_BLK_NAME}:{L2}{R_BLK_NAME}')

        ws[f'{L}{R_BLK_PERIOD}'] = (
            f'=IF($A${r}="","",IF(OR($B${r}="",$C${r}=""),"— період не заданий —",'
            f'TEXT(DAY($B${r}),"00")&"."&TEXT(MONTH($B${r}),"00")&"."&YEAR($B${r})&" – "&'
            f'TEXT(DAY($C${r}),"00")&"."&TEXT(MONTH($C${r}),"00")&"."&YEAR($C${r})))')
        ws[f'{L}{R_BLK_STATUS}'] = f'=IF($A${r}="","",$D${r})'
        ws[f'{L}{R_BLK_UPD}']    = f'=IF($A${r}="","",$E${r})'
        for rr in (R_BLK_PERIOD, R_BLK_STATUS, R_BLK_UPD):
            ws[f'{L}{rr}'].font = f(10, it=True, color='2A3B33')
            ws[f'{L}{rr}'].fill = fill(GREEN_PL); ws[f'{L2}{rr}'].fill = fill(GREEN_PL)
            ws[f'{L}{rr}'].alignment = Alignment(vertical='center', horizontal='center')
            ws.merge_cells(f'{L}{rr}:{L2}{rr}')

        # базові RO-ціни
        col_head(ws,f'{L}11','RO 1–2 ос.'); col_head(ws,f'{L2}11','Вихідні 1–2\n(авто)')
        for j in range(4):
            rr = R_BASE0 + j
            inp(ws, f'{L}{rr}', (src['base'][j] if src else None), MONEY)
            mrow = R_MULT0 + data['matcats'].index(data['basecats'][j])
            ws[f'{L2}{rr}'] = (f'=IF(OR({L}{rr}="",{L2}{mrow}=""),"",'
                               f'ROUND({L}{rr}*{L2}{mrow}/ОКРУГЛЕННЯ,0)*ОКРУГЛЕННЯ)')
            out(ws, f'{L2}{rr}', None, MONEY, band=True)

        # Δ варіант A — між опорними категоріями
        col_head(ws,f'{L}18','Δ грн'); col_head(ws,f'{L2}18','Δ %')
        DA_PAIRS = [(1,0),(2,0),(3,2)]
        for j,(a,b) in enumerate(DA_PAIRS):
            rr, ra, rb = R_DA0 + j, R_BASE0 + a, R_BASE0 + b
            ws[f'{L}{rr}']  = f'=IF(OR({L}{ra}="",{L}{rb}=""),"",{L}{ra}-{L}{rb})'
            ws[f'{L2}{rr}'] = f'=IF(OR({L}{ra}="",{L}{rb}="",{L}{rb}=0),"",{L}{ra}/{L}{rb}-1)'
            out(ws,f'{L}{rr}',None,MONEY,band=True); out(ws,f'{L2}{rr}',None,PCT,band=True)

        # Δ варіант B — до еталонного прайс-листа
        col_head(ws,f'{L}25','Δ грн'); col_head(ws,f'{L2}25','Δ %')
        for j in range(4):
            rr, rb = R_DB0 + j, R_BASE0 + j
            ref = (f'IFERROR(INDEX(БАЗА_RO,{j+1},IFERROR(MATCH(ЕТАЛОН_ПЛ,ПЛ_РЯДОК,0),0)),"")')
            ws[f'{L}{rr}']  = f'=IF(OR({L}{rb}="",ЕТАЛОН_ПЛ="",{ref}=""),"",{L}{rb}-{ref})'
            ws[f'{L2}{rr}'] = f'=IF(OR({L}{rb}="",ЕТАЛОН_ПЛ="",{ref}="",{ref}=0),"",{L}{rb}/{ref}-1)'
            out(ws,f'{L}{rr}',None,MONEY,band=True); out(ws,f'{L2}{rr}',None,PCT,band=True)

        # множники
        col_head(ws,f'{L}32','Множник ціни'); col_head(ws,f'{L2}32','Множник вихідних')
        for j in range(N_CAT):
            rr = R_MULT0 + j
            m = src['mult'][j] if src else (None,None)
            inp(ws,f'{L}{rr}', m[0], MULT); inp(ws,f'{L2}{rr}', m[1], MULT)

        # пакети
        col_head(ws,f'{L}53','Дорослий, грн'); col_head(ws,f'{L2}53','Дитина, грн')
        for j in range(4):
            rr = R_PKG0 + j
            p = src['pkg'][j] if src else (None,None)
            inp(ws,f'{L}{rr}', p[0], MONEY); inp(ws,f'{L2}{rr}', p[1], MONEY)

        # фіксований SPA
        col_head(ws,f'{L}60','Фіксовано, грн'); col_head(ws,f'{L2}60','—')
        for j in range(2):
            rr = R_SPA0 + j
            inp(ws,f'{L}{rr}', (src['spa'][j] if src else None), MONEY)
            ws[f'{L2}{rr}'].fill = fill(GREY_T); ws[f'{L2}{rr}'].border = BOX

        # перемикачі знижки — за замовчуванням 0 (поточна поведінка файлу)
        col_head(ws,f'{L}65','1 / 0'); col_head(ws,f'{L2}65','—')
        for j in range(3):
            rr = R_SW0 + j
            inp(ws,f'{L}{rr}', 0 if src else None, '0')
            ws[f'{L}{rr}'].alignment = Alignment(horizontal='center')
            ws[f'{L2}{rr}'].fill = fill(GREY_T); ws[f'{L2}{rr}'].border = BOX

    dv_sw = DataValidation(type='list', formula1='"0,1"', allow_blank=True, showErrorMessage=True,
                           errorTitle='Перемикач знижки', error='Допустимі значення: 0 або 1.')
    ws.add_data_validation(dv_sw); dv_sw.add(f'I{R_SW0}:{LASTL}{R_SW1}')
    dv_pos = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0',
                            allow_blank=True, showErrorMessage=True, errorTitle='Некоректне значення',
                            error='Ціни та множники не можуть бути відʼємними.')
    ws.add_data_validation(dv_pos)
    for rng in (f'I{R_BASE0}:{LASTL}{R_BASE1}', f'I{R_MULT0}:{LASTL}{R_MULT1}',
                f'I{R_PKG0}:{LASTL}{R_PKG1}', f'I{R_SPA0}:{LASTL}{R_SPA1}'):
        dv_pos.add(rng)
