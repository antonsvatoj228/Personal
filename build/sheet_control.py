from openpyxl.styles import Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.utils import get_column_letter as gl
from style import *
from names import *

R0 = 7
R1 = R0 + N_CAT - 1
HKEYS = ['ro','mult','wmult','maxg','incl','ming','pa_b','pk_b','pa_v','pk_v','spa','okd','oke']
H1_0, H2_0 = 23, 37                      # W..AI  та  AK..AW
N_TOT = '($D$4+$F$4)'

def hcols(start):
    return {k: gl(start + i) for i, k in enumerate(HKEYS)}

def build(wb, data):
    ws = wb.create_sheet('Контроль і порівняння')
    title(ws,'A1','ПОРІВНЯННЯ ТА КОНТРОЛЬ ПРАЙС-ЛИСТІВ — BUKOVEL','A1:I1')
    for w,c in zip([40,13,13,13,13,12,12,11,11,13,12,13,12,12,12],
                   ['A','B','C','D','E','F','G','H','I','J','K','L','M','N','O']):
        ws.column_dimensions[c].width = w
    ws.column_dimensions['P'].width = 3

    for addr,lab,val,fmt in (('B3','Прайс-лист 1','Осінь 2026 (00-10%)',None),
                             ('D3','Тариф 1','Rack Rate',None),
                             ('F3','Прайс-лист 2','Осінь 2026 (10-20%)',None),
                             ('H3','Тариф 2','Rack Rate',None),
                             ('B4','Пакет','RO',None),
                             ('D4','Дорослих',2,'0'),
                             ('F4','Дітей',0,'0')):
        lcell = gl(ws[addr].column - 1) + addr[1:]
        label(ws,lcell,lab,bold=True); inp(ws,addr,val,fmt)
    for rng,src in (('B3','=ПРАЙС_ЛИСТИ'),('F3','=ПРАЙС_ЛИСТИ'),('D3','=ТАРИФИ_НАЗВИ'),
                    ('H3','=ТАРИФИ_НАЗВИ'),('B4','=СПИСОК_ПАКЕТІВ')):
        dv = DataValidation(type='list', formula1=src, allow_blank=False)
        ws.add_data_validation(dv); dv.add(rng)
    dv_g = DataValidation(type='whole', operator='between', formula1='0', formula2='20',
                          allow_blank=False, showErrorMessage=True,
                          errorTitle='Кількість гостей', error='Введіть ціле число від 0 до 20.')
    ws.add_data_validation(dv_g); dv_g.add('D4'); dv_g.add('F4')


    for i,h in enumerate(['Категорія','Прайс 1\nбудні','Прайс 1\nвихідні','Прайс 2\nбудні','Прайс 2\nвихідні',
                          'Δ грн\nбудні','Δ грн\nвихідні','Δ %\nбудні','Δ %\nвихідні',
                          'Прайс 1\nбуд.→вих.\nгрн','Прайс 1\nбуд.→вих.\n%',
                          'Прайс 2\nбуд.→вих.\nгрн','Прайс 2\nбуд.→вих.\n%',
                          'Δ спреду\nгрн','Δ спреду\nп.п.']):
        col_head(ws,f'{gl(i+1)}6',h)
    for c0,c1,t in ((2,5,'ЦІНА ЗА НІЧ'),(6,9,'РІЗНИЦЯ МІЖ ПРАЙС-ЛИСТАМИ'),
                    (10,13,'РІЗНИЦЯ БУДНІ → ВИХІДНІ В МЕЖАХ ПРАЙС-ЛИСТА'),(14,15,'ПОРІВНЯННЯ СПРЕДУ')):
        block_head(ws,f'{gl(c0)}5',t,f'{gl(c0)}5:{gl(c1)}5')
        ws[f'{gl(c0)}5'].alignment = Alignment(horizontal='center',vertical='center')
    ws.row_dimensions[5].height = 20
    ws.row_dimensions[6].height = 34

    _globals(ws)
    _helpers(ws,'$B$3','$D$3',H1_0,'K')
    _helpers(ws,'$F$3','$H$3',H2_0,'L')

    for i in range(N_CAT):
        r, dr, band = R0+i, 5+i, i % 2 == 1
        ws[f'A{r}'] = f"=IF('Довідник категорій'!A{dr}=\"\",\"\",'Довідник категорій'!A{dr})"
        out(ws,f'A{r}',None,band=band); ws[f'A{r}'].font = f(b=True)
        ws[f'B{r}'] = _price(H1_0,'K',False,r); ws[f'C{r}'] = _price(H1_0,'K',True,r)
        ws[f'D{r}'] = _price(H2_0,'L',False,r); ws[f'E{r}'] = _price(H2_0,'L',True,r)
        ws[f'F{r}'] = f'=IF(OR(B{r}="",D{r}=""),"",D{r}-B{r})'
        ws[f'G{r}'] = f'=IF(OR(C{r}="",E{r}=""),"",E{r}-C{r})'
        ws[f'H{r}'] = f'=IF(OR(B{r}="",D{r}="",B{r}=0),"",D{r}/B{r}-1)'
        ws[f'I{r}'] = f'=IF(OR(C{r}="",E{r}="",C{r}=0),"",E{r}/C{r}-1)'
        ws[f'J{r}'] = f'=IF(OR(B{r}="",C{r}=""),"",C{r}-B{r})'
        ws[f'K{r}'] = f'=IF(OR(B{r}="",C{r}="",B{r}=0),"",C{r}/B{r}-1)'
        ws[f'L{r}'] = f'=IF(OR(D{r}="",E{r}=""),"",E{r}-D{r})'
        ws[f'M{r}'] = f'=IF(OR(D{r}="",E{r}="",D{r}=0),"",E{r}/D{r}-1)'
        ws[f'N{r}'] = f'=IF(OR(J{r}="",L{r}=""),"",L{r}-J{r})'
        ws[f'O{r}'] = f'=IF(OR(K{r}="",M{r}=""),"",M{r}-K{r})'
        for c in 'BCDEFGJLN': out(ws,f'{c}{r}',None,MONEY,band=band)
        for c in 'HIKMO':     out(ws,f'{c}{r}',None,PCT,band=band)

    ws.conditional_formatting.add(f'A{R0}:O{R1}',
        FormulaRule(formula=[f'IFERROR(INDEX(ДОВ_РОЛЬ,MATCH($A{R0},КАТЕГОРІЇ,0))="Базова",FALSE)'],
                    fill=fill(RED_SOFT), font=f(11,True,RED_INK), stopIfTrue=False))
    for rng in (f'F{R0}:G{R1}', f'H{R0}:I{R1}', f'N{R0}:N{R1}', f'O{R0}:O{R1}'):
        ws.conditional_formatting.add(rng, CellIsRule(operator='lessThan', formula=['0'],
                                                      font=f(11,color=RED_INK)))
    _checklist(ws)

    for c in range(21, H2_0 + len(HKEYS)):
        ws.column_dimensions[gl(c)].hidden = True
    ws.freeze_panes = f'B{R0}'
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GREEN_M
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f'A1:O{R1}'; ws.print_title_rows = '1:6'
    return ws


def _globals(ws):
    ws['U1'] = 'ТЕХНІЧНІ ДАНІ — НЕ РЕДАГУВАТИ'; ws['U1'].font = f(10,True,'6B7771')
    for r,fml,lab in ((3,"=IFERROR(MATCH('Прайс'!$B$3,ПЛ_РЯДОК,0),0)",'колонка активного прайсу'),
                      (5,'=IF(OR($B$4="BB",$B$4="BBSPA",$B$4="ROSPABB"),1,0)','потрібен пакет'),
                      (6,'=IF($U$5=0,"",IF($B$4="ROSPABB","BB",$B$4))','ключ пакета'),
                      (7,'=IF(OR($B$4="ROSPA",$B$4="ROSPABB"),1,0)','потрібен SPA')):
        ws[f'U{r}'] = fml; ws[f'U{r}'].font = f(9,color='6B7771')
        ws[f'V{r}'] = lab; ws[f'V{r}'].font = f(9,it=True,color='9AA6A0')


def _helpers(ws, pl_cell, tar_cell, start, tag):
    c = hcols(start); H = gl(start)
    ws[f'{H}1'] = f'{tag}: службові'; ws[f'{H}1'].font = f(9,True,'6B7771')
    ws[f'{H}2'] = f'=IFERROR(MATCH({pl_cell},ПЛ_РЯДОК,0),0)'
    ws[f'{H}3'] = (f'=IF($B$4="BBSPA",ЗНИЖКА_BBSPA,IFERROR(SUMIFS(ТАРИФИ_ЗНИЖКИ,ТАРИФИ_НАЗВИ,{tar_cell}),0))')
    ws[f'{H}4'] = f'=IF({H}$2=0,0,IFERROR(INDEX(ПЕРЕМИКАЧІ,IF($B$4="BBSPA",2,1),{H}$2),0))'
    ws[f'{H}5'] = f'=IF({H}$2=0,0,IFERROR(INDEX(ПЕРЕМИКАЧІ,3,{H}$2),0))'
    for r in range(2,6): ws[f'{H}{r}'].font = f(9,color='6B7771')
    for i in range(N_CAT):
        r = R0 + i
        m = f'MATCH($A{r},КАТЕГОРІЇ,0)'
        ws[f'{c["ro"]}{r}']   = (f'=IF(OR({H}$2=0,$A{r}=""),0,IFERROR(INDEX(БАЗА_RO,'
                                f'MATCH(INDEX(ДОВ_ОПОРА,{m}),БАЗА_КАТ,0),{H}$2),0))')
        ws[f'{c["mult"]}{r}'] = f'=IF(OR({H}$2=0,$A{r}=""),0,IFERROR(INDEX(МНОЖНИКИ,MATCH($A{r},МНОЖ_КАТ,0),{H}$2),0))'
        ws[f'{c["wmult"]}{r}']= f'=IF(OR({H}$2=0,$A{r}=""),0,IFERROR(INDEX(МНОЖНИКИ,MATCH($A{r},МНОЖ_КАТ,0),{H}$2+1),0))'
        ws[f'{c["maxg"]}{r}'] = f'=IF($A{r}="",0,IFERROR(INDEX(ДОВ_МАКС,{m}),0))'
        ws[f'{c["incl"]}{r}'] = f'=IF($A{r}="",0,IFERROR(INDEX(ДОВ_ВКЛ,{m}),0))'
        ws[f'{c["ming"]}{r}'] = f'=IF($A{r}="",0,IFERROR(INDEX(ДОВ_МІН,{m}),0))'
        for key,day,off in (('pa_b','Будні',0),('pk_b','Будні',1),('pa_v','Вихідні',0),('pk_v','Вихідні',1)):
            ws[f'{c[key]}{r}'] = (f'=IF(OR({H}$2=0,$U$5=0),0,IFERROR(INDEX(ПАКЕТИ_ЦІНИ,'
                                  f'MATCH($U$6&"|{day}",ПАКЕТИ_КЛЮЧ,0),{H}$2+{off}),0))')
        ws[f'{c["spa"]}{r}']  = (f'=IF(OR({H}$2=0,$U$7=0,$A{r}=""),0,IFERROR(INDEX(SPA_ЦІНИ,'
                                f'MATCH(INDEX(ДОВ_SPA,{m}),SPA_РОЛЬ,0),{H}$2),0))')
        ws[f'{c["okd"]}{r}']  = (f'=IF(OR($A{r}="",{c["ro"]}{r}=0,{c["mult"]}{r}=0,'
                                f'AND($U$5=1,$D$4>0,{c["pa_b"]}{r}=0),AND($U$5=1,$F$4>0,{c["pk_b"]}{r}=0),'
                                f'AND($U$7=1,{c["spa"]}{r}=0)),0,1)')
        ws[f'{c["oke"]}{r}']  = (f'=IF(OR({c["okd"]}{r}=0,{c["wmult"]}{r}=0,'
                                f'AND($U$5=1,$D$4>0,{c["pa_v"]}{r}=0),AND($U$5=1,$F$4>0,{c["pk_v"]}{r}=0)),0,1)')
        for k in c.values(): ws[f'{k}{r}'].font = f(9,color='6B7771')


def _price(start, tag, wknd, r):
    c = hcols(start); H = gl(start)
    ok  = c['oke'] if wknd else c['okd']
    pa  = c['pa_v'] if wknd else c['pa_b']
    pk  = c['pk_v'] if wknd else c['pk_b']
    core = f'{c["ro"]}{r}*{c["mult"]}{r}' + (f'*{c["wmult"]}{r}' if wknd else '')
    return (f'=IF(OR({ok}{r}=0,{N_TOT}=0,{N_TOT}>{c["maxg"]}{r},{N_TOT}<{c["ming"]}{r}),"",'
            f'ROUND({core}*(1-{H}$3)/ОКРУГЛЕННЯ,0)*ОКРУГЛЕННЯ'
            f'+MAX(0,{N_TOT}-{c["incl"]}{r})*КРОК_ГОСТЯ*(1-{H}$3)'
            f'+($D$4*{pa}{r}+$F$4*{pk}{r})*(1-{H}$3*{H}$4)'
            f'+{c["spa"]}{r}*(1-{H}$3*{H}$5))')


def _checklist(ws):
    block_head(ws,'Q3','КОНТРОЛЬ АКТИВНОГО ПРАЙС-ЛИСТА','Q3:S3')
    ws.column_dimensions['Q'].width = 46; ws.column_dimensions['R'].width = 15
    ws.column_dimensions['S'].width = 3
    checks = [
      (4,'Незаповнені базові RO-ціни',
       '=4-COUNT(INDEX(БАЗА_RO,1,$U$3),INDEX(БАЗА_RO,2,$U$3),INDEX(БАЗА_RO,3,$U$3),INDEX(БАЗА_RO,4,$U$3))'),
      (5,'Незаповнені множники',
       f'={2*N_CAT}-COUNT(INDEX(МНОЖНИКИ,0,$U$3))-COUNT(INDEX(МНОЖНИКИ,0,$U$3+1))'),
      (6,'Незаповнені пакети BB',
       '=4-COUNT(INDEX(ПАКЕТИ_ЦІНИ,1,$U$3),INDEX(ПАКЕТИ_ЦІНИ,1,$U$3+1),'
       'INDEX(ПАКЕТИ_ЦІНИ,2,$U$3),INDEX(ПАКЕТИ_ЦІНИ,2,$U$3+1))'),
      (7,'Незаповнений фіксований SPA-візит',
       '=2-COUNT(INDEX(SPA_ЦІНИ,1,$U$3),INDEX(SPA_ЦІНИ,2,$U$3))'),
      (8,'Дублікати назв прайс-листів',
       '=SUMPRODUCT((ПРАЙС_ЛИСТИ<>"")*(COUNTIF(ПРАЙС_ЛИСТИ,ПРАЙС_ЛИСТИ&"")>1))'),
    ]
    for r,lab,fml in checks:
        label(ws,f'Q{r}',lab); ws[f'R{r}'] = fml; out(ws,f'R{r}',None,'0')
        ws.conditional_formatting.add(f'R{r}', CellIsRule(operator='greaterThan', formula=['0'],
                                     fill=fill(RED_SOFT), font=f(11,True,RED_INK)))
    ws['Q10'] = ('=IF(SUM($R$4:$R$8)=0,"ГОТОВО ДО ПУБЛІКАЦІЇ","ПОТРІБНЕ ЗАПОВНЕННЯ")')
    ws['Q10'].font = f(12,True); ws['Q10'].alignment = Alignment(horizontal='center',vertical='center')
    ws.merge_cells('Q10:R10'); ws.row_dimensions[10].height = 26
    ws.conditional_formatting.add('Q10:R10', FormulaRule(formula=['$Q$10="ГОТОВО ДО ПУБЛІКАЦІЇ"'],
                                  fill=fill('E1EFE7'), font=f(12,True,'2E6B4F'), stopIfTrue=True))
    ws.conditional_formatting.add('Q10:R10', FormulaRule(formula=['$Q$10="ПОТРІБНЕ ЗАПОВНЕННЯ"'],
                                  fill=fill(RED_SOFT), font=f(12,True,RED_INK), stopIfTrue=True))
    info = [
      (12,'Пакет BBSPA заповнено?',
       '=IF(COUNT(INDEX(ПАКЕТИ_ЦІНИ,3,$U$3),INDEX(ПАКЕТИ_ЦІНИ,3,$U$3+1),'
       'INDEX(ПАКЕТИ_ЦІНИ,4,$U$3),INDEX(ПАКЕТИ_ЦІНИ,4,$U$3+1))=4,"так","ні — пакет недоступний")',None),
      (13,'Прайс-листів у моделі','=COUNTIF(ПРАЙС_ЛИСТИ,"<>")','0'),
      (14,'Категорій у моделі','=COUNTIF(КАТЕГОРІЇ,"<>")','0'),
      (15,'Крок за 3-го гостя, грн','=КРОК_ГОСТЯ',MONEY),
      (16,'Крок округлення, грн','=ОКРУГЛЕННЯ',MONEY),
    ]
    for r,lab,fml,fmt in info:
        label(ws,f'Q{r}',lab); ws[f'R{r}'] = fml; out(ws,f'R{r}',None,fmt,band=True)
    label(ws,'Q18','BBSPA не входить до чек-листа: пакет може бути свідомо не задіяний у сезоні. '
                   'Його стан показано окремим рядком вище.')
    ws['Q18'].font = f(9,it=True,color='6B7771')
    ws['Q18'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells('Q18:R19'); ws.row_dimensions[18].height = 26
