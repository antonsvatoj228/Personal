from openpyxl.styles import Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from style import *
from names_s import *

HEAD = ['Код','Категорія','Назва на Booking','Мін. гостей','Включено у базову','Макс. гостей',
        'Доплата: дорослий, грн','Доплата: дитина, грн',"SPA-рівень для ROSPA / ROSPABB",'Статус']
WID  = [10,28,30,12,17,13,17,17,30,10]

def build(wb, data):
    ws = wb.create_sheet('Довідник категорій')
    title(ws,'A1','ДОВІДНИК КАТЕГОРІЙ — PHOENIX СХІДНИЦЯ','A1:J1')
    banner(ws,'A2','У щоденній роботі цей лист не редагують. Він задає місткість, базову місткість (скільки гостей уже входить у ціну), '
                   'доплату за додаткове місце та SPA-рівень. Зміна «Макс. гостей» або «Включено у базову» одразу змінює сітку '
                   'у «Прайс», «Калькулятор» і «Контроль і порівняння». Вільні рядки нижче — під нові категорії.','A2:J2',42)
    for i,h in enumerate(HEAD):
        col_head(ws,f'{col_letter(i+1)}4',h)
        ws.column_dimensions[col_letter(i+1)].width = WID[i]
    ws.row_dimensions[4].height = 34

    for i in range(N_CAT):
        r = 5 + i
        src = data['cats'][i] if i < len(data['cats']) else None
        band = (i % 2 == 1)
        vals = (list(src) if src else [None]*9)
        # оригінал: код, категорія, booking, мін, вкл, макс, допл_дор, допл_діт, статус
        code, name, bk, mn, incl, mx, da, dk_, st = vals
        if name == 'Premium Chalet':
            mx = 8           # узгоджено: «7, 8 дорослий по додатковому місцю» (у файлі було 7)
        spa = None
        if name in ('Classic Chalet','Forest View Chalet'): spa = "Малий об'єкт"
        elif name == 'Premium Chalet':                       spa = 'Великий об’єкт'
        row = [code, name, bk, mn, incl, mx, da, dk_, spa, st or ('OK' if name else None)]
        for c,v in enumerate(row, start=1):
            cell = f'{col_letter(c)}{r}'
            if c in (4,5,6,7,8):
                inp(ws, cell, int(v) if isinstance(v,float) else v, '0' if c in (4,5,6) else MONEY)
                ws[cell].alignment = Alignment(horizontal='center')
            else:
                out(ws, cell, v, band=band)
        ws[f'B{r}'].font = f(b=True)

    dvs = [(DataValidation(type='list', formula1='=СПИСОК_SPA', allow_blank=True), f'I5:I{DR1}'),
           (DataValidation(type='list', formula1='"OK,Не продається"', allow_blank=True), f'J5:J{DR1}'),
           (DataValidation(type='whole', operator='between', formula1='0', formula2='30', allow_blank=True,
                           showErrorMessage=True, errorTitle='Місткість',
                           error='Введіть ціле число від 0 до 30.'), f'D5:F{DR1}'),
           (DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True,
                           showErrorMessage=True, errorTitle='Доплата',
                           error='Доплата не може бути відʼємною.'), f'G5:H{DR1}')]
    for dv,rng in dvs:
        ws.add_data_validation(dv); dv.add(rng)
    ws.freeze_panes = 'C5'
    ws.sheet_view.showGridLines = False
    return ws
