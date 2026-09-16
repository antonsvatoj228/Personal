from openpyxl.styles import Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from style import *
from names_s import *

def build(wb, data):
    ws = wb.create_sheet('Калькулятор', 1)
    title(ws,'A1','КАЛЬКУЛЯТОР ВАРТОСТІ БРОНЮВАННЯ — ПHOENIX СХІДНИЦЯ','A1:F1')
    ws['A1'] = 'КАЛЬКУЛЯТОР ВАРТОСТІ БРОНЮВАННЯ — PHOENIX СХІДНИЦЯ'
    banner(ws,'A3','Заповнюйте лише жовті клітинки. Рахує тією самою формулою, що й лист «Прайс», але додатково розділяє '
                   'дорослих і дітей: базову місткість заповнюють спершу дорослі, решта гостей іде як додаткові місця. '
                   'Поле «Дата заїзду» необовʼязкове — воно лише підказує прайс-лист і тип дня.','A3:F3',42)
    for w,c in zip([32,22,3,38,3,18],'ABCDEF'): ws.column_dimensions[c].width = w
    block_head(ws,'A5','ВХІДНІ ПАРАМЕТРИ','A5:B5'); block_head(ws,'D5','РОЗРАХУНОК','D5:F5')
    ws['A6'] = 'Активний прайс-лист'; ws['A6'].font = f(b=True)
    ws['B6'] = "='Прайс'!$B$3"; ws['B6'].fill = fill(GREEN_PL); ws['B6'].font = f(b=True,color=GREEN_D); ws['B6'].border = BOX
    for r,lab,val,fmt in ((7,'Тариф (знижка)','Rack Rate',None),(8,'Категорія','Classic Chalet',None),
                          (9,'Тип дня','Будні',None),(10,'Дорослих',2,'0'),(11,'Дітей',0,'0'),
                          (12,'Пакет','RO',None),(13,'Кількість тварин',0,'0'),(14,'Кількість ночей',1,'0'),
                          (15,'Дата заїзду (необовʼязково)',None,DATE)):
        label(ws,f'A{r}',lab,bold=True); inp(ws,f'B{r}',val,fmt)
    for rng,src in (('B7','=ТАРИФИ_НАЗВИ'),('B8','=КАТЕГОРІЇ'),('B9','=СПИСОК_ДНІВ'),('B12','=СПИСОК_ПАКЕТІВ')):
        dv = DataValidation(type='list',formula1=src,allow_blank=False,showErrorMessage=True,
                            errorTitle='Оберіть зі списку',error='Значення має бути з випадаючого списку.')
        ws.add_data_validation(dv); dv.add(rng)
    dvg = DataValidation(type='whole',operator='between',formula1='0',formula2='30',allow_blank=False,
                         showErrorMessage=True,errorTitle='Кількість',error='Введіть ціле число від 0 до 30.')
    ws.add_data_validation(dvg); dvg.add('B10:B11'); dvg.add('B13')
    dvn = DataValidation(type='whole',operator='between',formula1='1',formula2='365',allow_blank=False,
                         showErrorMessage=True,errorTitle='Кількість ночей',error='Введіть ціле число від 1 до 365.')
    ws.add_data_validation(dvn); dvn.add('B14')
    dvd = DataValidation(type='date',operator='between',formula1='DATE(2020,1,1)',formula2='DATE(2040,12,31)',
                         allow_blank=True,showErrorMessage=True,errorTitle='Некоректна дата',
                         error='Введіть справжню дату у діапазоні 2020–2040.')
    ws.add_data_validation(dvd); dvd.add('B15')

    M = 'MATCH($B$8,КАТЕГОРІЇ,0)'
    rows = [
      (6,'Статус',
       '=IF(OR($B$6="",$B$8="",$U$3=0,F7<F8,F7>F10,F11=0,'
       'AND($B$9="Вихідні",F25=0),'
       'AND($U$8=1,OR(AND($B$10>0,F18=0),AND($B$11>0,F19=0))),AND($U$9=1,F20=0)),'
       '"ПЕРЕВІРТЕ ПАРАМЕТРИ","OK")',None),
      (7,'Усього гостей','=$B$10+$B$11','0'),
      (8,'Мін. гостей',f'=IFERROR(INDEX(ДОВ_МІН,{M}),0)','0'),
      (9,'Включено у базову ціну',f'=IFERROR(INDEX(ДОВ_ВКЛ,{M}),0)','0'),
      (10,'Макс. гостей',f'=IFERROR(INDEX(ДОВ_МАКС,{M}),0)','0'),
      (11,'RO база буднього дня, грн',
       '=IF($U$3=0,0,IFERROR(INDEX(БАЗА_ЦІНИ,MATCH($B$8,БАЗА_КАТ,0),$U$3),0))',MONEY),
      (12,'Додаткових дорослих','=MAX(0,$B$10-F9)','0'),
      (13,'Додаткових дітей','=MAX(0,F7-F9)-F12','0'),
      (14,'Доплата за додаткові місця, грн',
       f'=F12*IFERROR(INDEX(ДОВ_ДОПЛ_ДОР,{M}),0)+F13*IFERROR(INDEX(ДОВ_ДОПЛ_ДІТ,{M}),0)',MONEY),
      (15,'Знижка на проживання','=IF($B$12="BBSPA",ЗНИЖКА_BBSPA,IFERROR(SUMIFS(ТАРИФИ_ЗНИЖКИ,ТАРИФИ_НАЗВИ,$B$7),0))','0%'),
      (16,'Перемикач: знижка на пакет','=IF($U$3=0,0,IFERROR(INDEX(ПЕРЕМИКАЧІ,IF($B$12="BBSPA",2,1),$U$3),0))','0'),
      (17,'Перемикач: знижка на SPA','=IF($U$3=0,0,IFERROR(INDEX(ПЕРЕМИКАЧІ,3,$U$3),0))','0'),
      (18,'Пакет, дорослий, грн',
       '=IF(OR($U$3=0,$U$8=0),0,IFERROR(INDEX(ПАКЕТИ_ЦІНИ,MATCH($U$10&"|"&$B$9,ПАКЕТИ_КЛЮЧ,0),$U$3),0))',MONEY),
      (19,'Пакет, дитина, грн',
       '=IF(OR($U$3=0,$U$8=0),0,IFERROR(INDEX(ПАКЕТИ_ЦІНИ,MATCH($U$10&"|"&$B$9,ПАКЕТИ_КЛЮЧ,0),$U$3+1),0))',MONEY),
      (20,'Фіксований SPA-візит, грн',
       f'=IF(OR($U$3=0,$U$9=0),0,IFERROR(INDEX(SPA_ЦІНИ,MATCH(INDEX(ДОВ_SPA,{M}),SPA_РІВЕНЬ,0),'
       '$U$3+IF($B$9="Вихідні",1,0)),0))',MONEY),
      (21,'Проживання з тваринами, грн','=$B$13*ТВАРИНИ',MONEY),
      (22,'Проживання за базовою місткістю до знижки, грн','=F11*IF($B$9="Вихідні",F25,1)','#,##0.00'),
      (23,'ЦІНА ЗА НІЧ, грн',
       '=IF($F$6<>"OK","",ROUND(F22*(1-F15)/ОКРУГЛЕННЯ,0)*ОКРУГЛЕННЯ'
       '+F14*(1-F15)+($B$10*F18+$B$11*F19)*(1-F15*F16)+F20*(1-F15*F17)+F21)',MONEY),
      (24,'ВАРТІСТЬ ПРОЖИВАННЯ, грн','=IF(F23="","",F23*$B$14)',MONEY),
    ]
    for r,lab,fml,fmt in rows:
        label(ws,f'D{r}',lab,wrap=True); ws[f'F{r}'] = fml; out(ws,f'F{r}',None,fmt)
        if r in (23,24):
            ws[f'D{r}'].font = f(12,True,GREEN_D)
            ws[f'F{r}'].font = f(12,True,GREEN_D); ws[f'F{r}'].fill = fill(GREEN_PL)
            ws.row_dimensions[r].height = 22
    ws.conditional_formatting.add('F6', FormulaRule(formula=['$F$6="OK"'],
        fill=fill('E1EFE7'), font=f(11,True,'2E6B4F'), stopIfTrue=True))
    ws.conditional_formatting.add('F6', FormulaRule(formula=['$F$6<>"OK"'],
        fill=fill(RED_SOFT), font=f(11,True,RED_INK), stopIfTrue=True))

    label(ws,'D25','Множник вихідних',wrap=True)
    ws['F25'] = '=IF($U$3=0,0,IFERROR(INDEX(МНОЖ_ВИХ,MATCH($B$8,МНОЖ_КАТ,0),$U$3),0))'
    out(ws,'F25',None,MULT)
    block_head(ws,'D27','ПІДКАЗКА ЗА ДАТОЮ ЗАЇЗДУ','D27:F27')
    label(ws,'D28','Прайс-лист, що покриває дату',wrap=True)
    ws['F28'] = ('=IF($B$15="","",IFERROR(INDEX(ПРАЙС_ЛИСТИ,SUMPRODUCT(MAX((ПЛ_ДАТА_З<=$B$15)*(ПЛ_ДАТА_ПО>=$B$15)'
                 '*(ПЛ_СТАТУС="Активний")*(ПРАЙС_ЛИСТИ<>"")*ROW(ПРАЙС_ЛИСТИ)))-' + str(R_PL0-1) + '),"— не знайдено —"))')
    label(ws,'D29','Тип дня за датою',wrap=True)
    ws['F29'] = '=IF($B$15="","",IF(WEEKDAY($B$15,2)>=ПЕРШИЙ_ВИХІДНИЙ,"Вихідні","Будні"))'
    for r in (28,29): out(ws,f'F{r}',None,band=True)
    label(ws,'D30','Підказка не змінює розрахунок — прайс-лист і тип дня обираються вручну вище.')
    ws['D30'].font = f(9,it=True,color='6B7771'); ws.merge_cells('D30:F30')

    for r,fml,lab in ((3,'=IFERROR(MATCH($B$6,ПЛ_РЯДОК,0),0)','колонка прайс-листа'),
                      (8,'=IF(OR($B$12="BB",$B$12="BBSPA",$B$12="ROSPABB"),1,0)','потрібен пакет'),
                      (9,'=IF(OR($B$12="ROSPA",$B$12="ROSPABB"),1,0)','потрібен SPA'),
                      (10,'=IF($U$8=0,"",IF($B$12="ROSPABB","BB",$B$12))','ключ пакета')):
        ws[f'U{r}'] = fml; ws[f'U{r}'].font = f(9,color='6B7771')
        ws[f'V{r}'] = lab; ws[f'V{r}'].font = f(9,it=True,color='9AA6A0')
    ws['U1'] = 'ТЕХНІЧНІ ДАНІ — НЕ РЕДАГУВАТИ'; ws['U1'].font = f(10,True,'6B7771')
    for c in ('U','V'): ws.column_dimensions[c].hidden = True
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GREEN_M
    ws.print_area = 'A1:F30'
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws
