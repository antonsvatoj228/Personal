from openpyxl.styles import Alignment, Border
from openpyxl.worksheet.datavalidation import DataValidation
from style import *
from names_s import *

PKG_LIST  = ['RO','BB','BBSPA','ROSPA','ROSPABB']
DAY_LIST  = ['Будні','Вихідні']
STAT_LIST = ['Активний','Неактивний','Архів']
SPA_LEVELS = ["Малий об'єкт", 'Великий об’єкт']
SW_LABELS = ['Знижка на пакет BB (сніданки)',
             'Знижка на пакет BBSPA',
             'Знижка на фіксований SPA-візит (ROSPA / ROSPABB)']

def build(wb, data):
    ws = wb.create_sheet('Керування прайсом')
    title(ws,'A1','КЕРУВАННЯ ПРАЙСОМ — PHOENIX СХІДНИЦЯ','A1:G1')
    banner(ws,'A2','Три кроки: 1) додайте прайс-лист у наступному вільному рядку зліва; 2) заповніть його двоколонковий блок праворуч '
                   '(будні / вихідні); 3) перевірте результат у листі «Прайс». Жовті клітинки — ввід.','A2:G2',34)

    # ── 1. Прайс-листи ──
    label(ws,'A4','1. ПРАЙС-ЛИСТИ / ПЕРІОДИ',bold=True); ws['A4'].font = f(11,True,GREEN_D)
    for c,h in zip('ABCDE',['Прайс-лист','Дата з','Дата по','Статус','Примітка / оновлено']):
        col_head(ws,f'{c}5',h)
    ws.row_dimensions[5].height = 30
    for w,c in zip([34,12,12,13,26],'ABCDE'): ws.column_dimensions[c].width = w
    for i in range(N_PL):
        r = R_PL0 + i
        src = data['pricelists'][i] if i < len(data['pricelists']) else None
        inp(ws,f'A{r}', src['name'] if src else None)
        inp(ws,f'B{r}', src.get('date_from') if src else None, DATE)
        inp(ws,f'C{r}', src.get('date_to')   if src else None, DATE)
        inp(ws,f'D{r}', src['status'] if src else None)
        inp(ws,f'E{r}', src.get('period_text') if src else None)
        ws[f'D{r}'].alignment = Alignment(horizontal='center')
    dv = DataValidation(type='list', formula1='=СПИСОК_СТАТУСІВ', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f'D{R_PL0}:D{R_PL_LAST}')
    dvd = DataValidation(type='date', operator='between', formula1='DATE(2020,1,1)', formula2='DATE(2040,12,31)',
                         allow_blank=True, showErrorMessage=True, errorTitle='Некоректна дата',
                         error='Введіть справжню дату у діапазоні 2020–2040.')
    ws.add_data_validation(dvd); dvd.add(f'B{R_PL0}:C{R_PL_LAST}')

    # ── 2. Глобальні параметри ──
    label(ws,'F4','2. ГЛОБАЛЬНІ ПАРАМЕТРИ',bold=True); ws['F4'].font = f(11,True,GREEN_D)
    ws.column_dimensions['F'].width = 46; ws.column_dimensions['G'].width = 24
    g = data['globals']
    glob = [('F5','Активний прайс-лист (показується у «Прайс»)',"='Прайс'!$B$3",None,False),
            ('F6','Крок округлення, грн',                        g['округлення'],  MONEY,True),
            ('F7','Знижка на проживання для пакета BBSPA',       g['знижка_bbspa'],'0%', True),
            ('F8','Проживання з тваринами, грн / ніч',           g['тварини'],     MONEY,True),
            ('F9','Еталонний прайс-лист для Δ (варіант B)',      data['pricelists'][1]['name'] if len(data['pricelists'])>1 else None, None, True),
            ('F10','Вихідні починаються з дня тижня (1=Пн … 7=Нд)', 5,             '0',  True)]
    for addr,lab,val,fmt,is_in in glob:
        label(ws,addr,lab,wrap=True); gc = 'G'+addr[1:]
        if is_in: inp(ws,gc,val,fmt)
        else:     out(ws,gc,val,fmt,bold=True)
    ws['G5'].fill = fill(GREEN_PL); ws['G5'].font = f(b=True,color=GREEN_D)
    label(ws,'F11','Доплата за додаткове місце задається по кожній категорії у «Довіднику категорій».')
    ws['F11'].font = f(9,it=True,color='6B7771'); ws.merge_cells('F11:G11')
    for dv2,rng in ((DataValidation(type='list',formula1='=ПРАЙС_ЛИСТИ',allow_blank=True),'G9'),
                    (DataValidation(type='decimal',operator='greaterThanOrEqual',formula1='0.01',allow_blank=False,
                        showErrorMessage=True,errorTitle='Крок округлення',
                        error='Крок округлення має бути не менший за 0,01 грн.'),'G6'),
                    (DataValidation(type='whole',operator='between',formula1='1',formula2='7',allow_blank=False,
                        showErrorMessage=True,errorTitle='День тижня',
                        error='Введіть число від 1 (понеділок) до 7 (неділя).'),'G10')):
        ws.add_data_validation(dv2); dv2.add(rng)

    # ── 3. Тарифи ──
    label(ws,'F12','3. ТАРИФИ — ЗНИЖКИ НА ПРОЖИВАННЯ',bold=True); ws['F12'].font = f(11,True,GREEN_D)
    col_head(ws,f'F{R_TAR0-1}','Тариф'); col_head(ws,f'G{R_TAR0-1}','Знижка')
    for i in range(N_TAR):
        r = R_TAR0 + i
        src = data['tariffs'][i] if i < len(data['tariffs']) else (None,None)
        inp(ws,f'F{r}',src[0]); inp(ws,f'G{r}',src[1],'0%')

    # ── 4. Довідники списків ──
    label(ws,'F23','4. ДОВІДНИКИ ДЛЯ ВИПАДАЮЧИХ СПИСКІВ',bold=True); ws['F23'].font = f(11,True,GREEN_D)
    col_head(ws,'F24','Пакети')
    for i,v in enumerate(PKG_LIST):  out(ws,f'F{25+i}',v)
    col_head(ws,'F31','Типи дня')
    for i,v in enumerate(DAY_LIST):  out(ws,f'F{32+i}',v)
    col_head(ws,'F35','Статуси')
    for i,v in enumerate(STAT_LIST): out(ws,f'F{36+i}',v)

    # ── Підписи рядків блоків ──
    ws.column_dimensions['H'].width = 46
    block_head(ws,f'H{R_BLK_NAME}','ПРАЙС-ЛИСТ')
    for r,t in ((R_BLK_PERIOD,'Період дії'),(R_BLK_STATUS,'Статус'),(R_BLK_UPD,'Оновлено')):
        label(ws,f'H{r}',t)
    block_head(ws,f'H{SECTION_ROWS[0]}','БАЗОВІ RO-ЦІНИ — ЗА БАЗОВОЮ МІСТКІСТЮ КАТЕГОРІЇ')
    col_head(ws,f'H{R_BASE0-1}','Категорія')
    _cat_labels(ws, R_BASE0)
    block_head(ws,f'H{SECTION_ROWS[1]}','Δ МІЖ КАТЕГОРІЯМИ  (варіант A)')
    col_head(ws,f'H{R_DA0-1}','Порівняння з попередньою категорією')
    _cat_labels(ws, R_DA0, prev=True)
    block_head(ws,f'H{SECTION_ROWS[2]}','Δ ДО ЕТАЛОННОГО ПРАЙС-ЛИСТА  (варіант B)')
    ws[f'H{R_DB0-2}'] = '=IF(ЕТАЛОН_ПЛ="","еталон не обрано — задайте G9","Еталон: "&ЕТАЛОН_ПЛ)'
    ws[f'H{R_DB0-2}'].font = f(10,it=True,color=GREEN_D); ws[f'H{R_DB0-2}'].fill = fill(GREEN_PL)
    col_head(ws,f'H{R_DB0-1}','Категорія')
    _cat_labels(ws, R_DB0)
    block_head(ws,f'H{SECTION_ROWS[3]}','ПАКЕТИ BB / BBSPA — ЗА 1 ГОСТЯ')
    col_head(ws,f'H{R_PKG0-1}','Пакет / тип дня')
    for i,n in enumerate(data['pkgkeys']): out(ws,f'H{R_PKG0+i}',n,bold=True)
    block_head(ws,f'H{SECTION_ROWS[4]}',"ФІКСОВАНИЙ SPA-ВІЗИТ ДЛЯ ROSPA / ROSPABB — 1 ОД. НА ОБ'ЄКТ")
    col_head(ws,f'H{R_SPA0-1}','SPA-рівень')
    for i,n in enumerate(SPA_LEVELS): out(ws,f'H{R_SPA0+i}',n,bold=True)
    block_head(ws,f'H{SECTION_ROWS[5]}','ПЕРЕМИКАЧІ ЗНИЖКИ   (1 = знижка тарифу діє, 0 = не діє)')
    col_head(ws,f'H{R_SW0-1}','На що діє знижка тарифу')
    for i,n in enumerate(SW_LABELS): out(ws,f'H{R_SW0+i}',n)
    label(ws,f'H{GUTTER_ROWS[-1]+1}',
          'Логіка: для кожної категорії задається власна RO-ціна окремо на будні та вихідні — множників немає, '
          'бо всі категорії базові. Крок округлення (G6) застосовується до базової ціни після знижки; доплата за кожне '
          'додаткове місце понад базову місткість береться з «Довідника категорій», множиться на знижку і не округлюється. '
          'BB / BBSPA додають пакет за кожного гостя. ROSPA додає один фіксований SPA-візит на обʼєкт, '
          'ROSPABB — сніданки за кількістю гостей + один SPA-візит.', wrap=True)
    nr = GUTTER_ROWS[-1]+1
    ws[f'H{nr}'].font = f(10,it=True)
    ws.merge_cells(f'H{nr}:N{nr}'); ws.row_dimensions[nr].height = 58

    _blocks(ws, data)
    _readability(ws)
    ws.freeze_panes = 'I5'
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GREEN_D
    ws.page_setup.orientation = 'landscape'
    return ws


def _cat_labels(ws, row0, prev=False):
    """Підписи категорій у колонці H — формулами з «Довідника», щоб нові категорії підхоплювались самі."""
    for i in range(N_CAT):
        r, dr = row0 + i, 5 + i
        if prev:
            ws[f'H{r}'] = (f'=IF(OR(\'Довідник категорій\'!B{dr}="",\'Довідник категорій\'!B{dr-1}=""),"",'
                           f'\'Довідник категорій\'!B{dr}&"  ↔  "&\'Довідник категорій\'!B{dr-1})'
                           ) if i > 0 else '=""'
        else:
            ws[f'H{r}'] = f'=IF(\'Довідник категорій\'!B{dr}="","",\'Довідник категорій\'!B{dr})'
        out(ws, f'H{r}', None, band=(i % 2 == 1))
        ws[f'H{r}'].font = f(10 if prev else 11, b=not prev)


def _blocks(ws, data):
    for i in range(N_PL):
        c = pl_col(i); L, L2 = col_letter(c), col_letter(c+1)
        r = R_PL0 + i
        src = data['pricelists'][i] if i < len(data['pricelists']) else None
        ws.column_dimensions[L].width = 15; ws.column_dimensions[L2].width = 15
        for cc in (L, L2):
            ws[f'{cc}{R_TECH}'] = f'=IF($A${r}="","",$A${r})'
            ws[f'{cc}{R_TECH}'].font = f(8,color='9AA6A0'); ws[f'{cc}{R_TECH}'].fill = fill(GREY_T)
        ws[f'{L}{R_BLK_NAME}'] = f'=IF($A${r}="","",$A${r})'
        ws[f'{L}{R_BLK_NAME}'].font = f(11,True,'FFFFFF'); ws[f'{L}{R_BLK_NAME}'].fill = fill(GREEN_D)
        ws[f'{L}{R_BLK_NAME}'].alignment = Alignment(vertical='center',wrap_text=True,indent=1)
        ws[f'{L2}{R_BLK_NAME}'].fill = fill(GREEN_D)
        ws.merge_cells(f'{L}{R_BLK_NAME}:{L2}{R_BLK_NAME}')
        ws[f'{L}{R_BLK_PERIOD}'] = (
            f'=IF($A${r}="","",IF(OR($B${r}="",$C${r}=""),IF($E${r}="","— період не заданий —",$E${r}),'
            f'TEXT(DAY($B${r}),"00")&"."&TEXT(MONTH($B${r}),"00")&"."&YEAR($B${r})&" – "&'
            f'TEXT(DAY($C${r}),"00")&"."&TEXT(MONTH($C${r}),"00")&"."&YEAR($C${r})))')
        ws[f'{L}{R_BLK_STATUS}'] = f'=IF($A${r}="","",$D${r})'
        ws[f'{L}{R_BLK_UPD}']    = f'=IF($A${r}="","",$E${r})'
        for rr in (R_BLK_PERIOD,R_BLK_STATUS,R_BLK_UPD):
            ws[f'{L}{rr}'].font = f(10,it=True,color='2A3B33')
            ws[f'{L}{rr}'].fill = fill(GREEN_PL); ws[f'{L2}{rr}'].fill = fill(GREEN_PL)
            ws[f'{L}{rr}'].alignment = Alignment(vertical='center',horizontal='center')
            ws.merge_cells(f'{L}{rr}:{L2}{rr}')

        # базові RO-ціни
        col_head(ws,f'{L}{R_BASE0-1}','RO будні'); col_head(ws,f'{L2}{R_BASE0-1}','RO вихідні')
        for j in range(N_CAT):
            rr = R_BASE0 + j
            b = src['base'][j] if (src and j < len(src['base'])) else (None,None)
            inp(ws,f'{L}{rr}', b[0], MONEY); inp(ws,f'{L2}{rr}', b[1], MONEY)

        # Δ між категоріями
        col_head(ws,f'{L}{R_DA0-1}','Δ грн'); col_head(ws,f'{L2}{R_DA0-1}','Δ %')
        for j in range(N_CAT):
            rr, cur, prev = R_DA0 + j, R_BASE0 + j, R_BASE0 + j - 1
            if j == 0:
                out(ws,f'{L}{rr}',None,MONEY,band=True); out(ws,f'{L2}{rr}',None,PCT,band=True); continue
            ws[f'{L}{rr}']  = f'=IF(OR({L}{cur}="",{L}{prev}=""),"",{L}{cur}-{L}{prev})'
            ws[f'{L2}{rr}'] = f'=IF(OR({L}{cur}="",{L}{prev}="",{L}{prev}=0),"",{L}{cur}/{L}{prev}-1)'
            out(ws,f'{L}{rr}',None,MONEY,band=True); out(ws,f'{L2}{rr}',None,PCT,band=True)

        # Δ до еталона
        col_head(ws,f'{L}{R_DB0-1}','Δ грн'); col_head(ws,f'{L2}{R_DB0-1}','Δ %')
        for j in range(N_CAT):
            rr, cur = R_DB0 + j, R_BASE0 + j
            ref = f'IFERROR(INDEX(БАЗА_ЦІНИ,{j+1},IFERROR(MATCH(ЕТАЛОН_ПЛ,ПЛ_РЯДОК,0),0)),"")'
            ws[f'{L}{rr}']  = f'=IF(OR({L}{cur}="",ЕТАЛОН_ПЛ="",{ref}=""),"",{L}{cur}-{ref})'
            ws[f'{L2}{rr}'] = f'=IF(OR({L}{cur}="",ЕТАЛОН_ПЛ="",{ref}="",{ref}=0),"",{L}{cur}/{ref}-1)'
            out(ws,f'{L}{rr}',None,MONEY,band=True); out(ws,f'{L2}{rr}',None,PCT,band=True)

        # пакети
        col_head(ws,f'{L}{R_PKG0-1}','Дорослий, грн'); col_head(ws,f'{L2}{R_PKG0-1}','Дитина, грн')
        for j in range(4):
            rr = R_PKG0 + j
            p = src['pkg'][j] if (src and j < len(src['pkg'])) else (None,None)
            inp(ws,f'{L}{rr}',p[0],MONEY); inp(ws,f'{L2}{rr}',p[1],MONEY)

        # фіксований SPA — будні / вихідні
        col_head(ws,f'{L}{R_SPA0-1}','Будні, грн'); col_head(ws,f'{L2}{R_SPA0-1}','Вихідні, грн')
        for j in range(2):
            rr = R_SPA0 + j
            inp(ws,f'{L}{rr}',None,MONEY); inp(ws,f'{L2}{rr}',None,MONEY)

        # перемикачі
        col_head(ws,f'{L}{R_SW0-1}','1 / 0'); col_head(ws,f'{L2}{R_SW0-1}','—')
        for j in range(3):
            rr = R_SW0 + j
            inp(ws,f'{L}{rr}', 0 if src else None, '0')
            ws[f'{L}{rr}'].alignment = Alignment(horizontal='center')
            ws[f'{L2}{rr}'].fill = fill(GREY_T); ws[f'{L2}{rr}'].border = BOX

    dvsw = DataValidation(type='list',formula1='"0,1"',allow_blank=True,showErrorMessage=True,
                          errorTitle='Перемикач знижки',error='Допустимі значення: 0 або 1.')
    ws.add_data_validation(dvsw); dvsw.add(f'I{R_SW0}:{LASTL}{R_SW1}')
    dvp = DataValidation(type='decimal',operator='greaterThanOrEqual',formula1='0',allow_blank=True,
                         showErrorMessage=True,errorTitle='Некоректне значення',
                         error='Ціни не можуть бути відʼємними.')
    ws.add_data_validation(dvp)
    for rng in (f'I{R_BASE0}:{LASTL}{R_BASE1}', f'I{R_PKG0}:{LASTL}{R_PKG1}', f'I{R_SPA0}:{LASTL}{R_SPA1}'):
        dvp.add(rng)


def _readability(ws):
    last = LAST_COL
    ws.row_dimensions[R_TECH].height = 10
    label(ws,f'H{R_TECH}','технічний рядок — не редагувати'); ws[f'H{R_TECH}'].font = f(8,it=True,color='9AA6A0')
    for r in SECTION_ROWS:
        for c in range(9,last+1):
            cell = ws.cell(row=r,column=c)
            cell.fill = fill(GREEN_M); cell.border = Border(left=THIN,right=THIN,top=MED,bottom=MED)
        ws.row_dimensions[r].height = 20
    for r in GUTTER_ROWS: gutter(ws, r, 8, last, height=9)
    for i in range(N_PL):
        c = pl_col(i)
        for r in list(range(2,9)) + list(range(10, GUTTER_ROWS[-1])):
            edge(ws, f'{col_letter(c)}{r}', left=True)
    for i in range(N_PL):
        r = R_PL0 + i
        if i % 2 == 1:
            for c in 'ABCDE': ws[f'{c}{r}'].fill = fill(INPUT_ALT)
        if (i+1) % 10 == 0:
            for c in 'ABCDE': edge(ws, f'{c}{r}', bottom=True)
    for r in range(R_PL0, R_PL_LAST+1): edge(ws, f'E{r}', right=True)
    edge(ws,'E5',right=True)
