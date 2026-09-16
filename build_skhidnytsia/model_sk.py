"""Незалежна модель Східниці."""
import openpyxl
from decimal import Decimal, ROUND_HALF_UP
def xlround(x, step):
    return float(Decimal(repr(x/step)).quantize(Decimal(1), rounding=ROUND_HALF_UP)) * step

def load(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    km, dov = wb['Керування прайсом'], wb['Довідник категорій']
    G = {'STEP': km['G6'].value, 'BBSPA': km['G7'].value, 'PETS': km['G8'].value}
    tariff = {km.cell(row=r,column=6).value: km.cell(row=r,column=7).value
              for r in range(14,22) if km.cell(row=r,column=6).value}
    plcol = {}
    for c in range(9, 209, 2):
        n = km.cell(row=2,column=c).value
        if n: plcol.setdefault(n, c)
    basecat = {km.cell(row=r,column=8).value: i for i,r in enumerate(range(12,24)) if km.cell(row=r,column=8).value}
    baserow = {km.cell(row=r,column=8).value: r for r in range(12,24) if km.cell(row=r,column=8).value}
    pkgr = {km.cell(row=r,column=8).value: r for r in range(58,62)}
    spar = {km.cell(row=r,column=8).value: r for r in range(65,67)}
    swr  = {'BB':70,'BBSPA':71,'SPA':72}
    cat = {}
    for r in range(5,17):
        nm = dov.cell(row=r,column=2).value
        if not nm: continue
        cat[nm] = dict(mn=dov.cell(row=r,column=4).value, incl=dov.cell(row=r,column=5).value,
                       mx=dov.cell(row=r,column=6).value, da=dov.cell(row=r,column=7).value,
                       dk=dov.cell(row=r,column=8).value, spa=dov.cell(row=r,column=9).value)
    def val(r,c):
        v = km.cell(row=r,column=c).value
        return 0 if v in (None,'') else v
    return dict(km=km, G=G, tariff=tariff, plcol=plcol, baserow=baserow,
                pkgr=pkgr, spar=spar, swr=swr, cat=cat, val=val)

def price(M, catn, n, pl, pkg, day, tar, adults=None, kids=0):
    if adults is None: adults, kids = n, 0
    G, val, cat = M['G'], M['val'], M['cat']
    if pl not in M['plcol'] or catn not in cat: return None
    c0, info = M['plcol'][pl], cat[catn]
    off = 1 if day == 'Вихідні' else 0
    base = val(M['baserow'][catn], c0+off)
    if n > info['mx'] or base == 0: return None
    need_pkg = pkg in ('BB','BBSPA','ROSPABB'); need_spa = pkg in ('ROSPA','ROSPABB')
    pa = pk = 0
    if need_pkg:
        key = ('BB|' if pkg=='ROSPABB' else pkg+'|') + day
        pa = val(M['pkgr'][key], c0); pk = val(M['pkgr'][key], c0+1)
        if pa == 0: return None
    spa = 0
    if need_spa:
        spa = val(M['spar'][info['spa']], c0+off)
        if spa == 0: return None
    disc = G['BBSPA'] if pkg == 'BBSPA' else M['tariff'].get(tar, 0)
    sw_pkg = val(M['swr']['BBSPA' if pkg=='BBSPA' else 'BB'], c0)
    sw_spa = val(M['swr']['SPA'], c0)
    ea = max(0, adults - info['incl'])
    ek = max(0, n - info['incl']) - ea
    return (xlround(base*(1-disc), G['STEP'])
            + (ea*info['da'] + ek*info['dk'])*(1-disc)
            + (adults*pa + kids*pk)*(1-disc*sw_pkg)
            + spa*(1-disc*sw_spa))
