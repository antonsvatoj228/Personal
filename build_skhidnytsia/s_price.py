from openpyxl.styles import Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as gl
from style import *
from names_s import *

R_GRP, R_HEAD, R0 = 7, 8, 9
R1 = R0 + N_CAT - 1
WD0, WD1 = 4, 4 + MAXG - 1          # D..K
SPACER   = WD1 + 1                  # L
WE0, WE1 = SPACER + 1, SPACER + MAXG  # M..T
HELP0 = WE1 + 1                     # U

PKG_DESC = ('=IF($B$5="RO","Лише проживання.",'
            'IF($B$5="BB","Проживання + сніданок за кожного гостя.",'
            'IF($B$5="BBSPA","Проживання + сніданок і SPA за кожного гостя.",'
            'IF($B$5="ROSPA","Проживання + 1 фіксований SPA-візит на обʼєкт.",'
            'IF($B$5="ROSPABB","Проживання + сніданки за кількістю гостей + 1 SPA-візит на обʼєкт.","")))))')

H = dict(ro_b='W', wm='X', ro_v='Y', maxg='Z', incl='AA', dopl='AB',
         pkg_b='AC', pkg_v='AD', spa_b='AE', spa_v='AF', ok_b='AG', ok_v='AH')

def build(wb, data):
    ws = wb.create_sheet('Прайс', 0)
    title(ws,'A1',None,f'A1:{gl(WE1)}1')
    ws['A1'] = '="ПРАЙС PHOENIX СХІДНИЦЯ · "&$B$3&" · "&$U$11&" · Тариф: "&$B$4&" · Пакет: "&$B$5'
    ws.row_dimensions[2].height = 6
    for r,(lab,dflt) in enumerate([('Активний прайс-лист','0сінь 26 0-10%'),
                                   ('Тариф (знижка)','Rack Rate'),
                                   ('Пакет','RO')], start=3):
        label(ws,f'A{r}',lab,bold=True); inp(ws,f'B{r}',dflt); ws.row_dimensions[r].height = 19
    ws['D3'] = '="Період дії: "&$U$11'
    ws['D4'] = ('="Знижка на проживання: "&TEXT($U$5,"0%")'
                '&IF($B$5="BBSPA"," — фіксована для пакета BBSPA; обраний тариф не застосовується","")')
    ws['D5'] = PKG_DESC
    for r in (3,4,5):
        ws[f'D{r}'].font = f(10,it=True,color='2A3B33'); ws[f'D{r}'].fill = fill(GREEN_PL)
        ws[f'D{r}'].alignment = Alignment(vertical='center',indent=1)
        for c in range(5, WE1+1): ws.cell(row=r,column=c).fill = fill(GREEN_PL)
        ws.merge_cells(start_row=r,start_column=4,end_row=r,end_column=WE1)

    banner(ws,'A6',None,f'A6:{gl(WE1)}6',30)
    ws['A6'] = ('=IF($U$3=0,"⚠ Прайс-лист не обрано або не знайдено — оберіть значення у клітинці B3.",'
                'IF($U$13=0,"⚠ У цьому прайс-листі ще не заповнені базові RO-ціни — сітка буде порожньою. '
                'Заповніть його блок на листі «Керування прайсом».",'
                '"Усі ціни — за ніч, за дорослих гостей, у гривні. Дитячі тарифи рахує лист «Калькулятор». '
                'Ціна діє за базовою місткістю категорії; кожне додаткове місце додає доплату з «Довідника категорій» × знижку."))')
    ws.conditional_formatting.add(f'A6:{gl(WE1)}6',
        FormulaRule(formula=['OR($U$3=0,$U$13=0)'], fill=fill(RED_SOFT), font=f(10,True,RED_INK), stopIfTrue=True))

    for c0,c1,t in ((WD0,WD1,'БУДНІ — кількість гостей'),(WE0,WE1,'ВИХІДНІ — кількість гостей')):
        block_head(ws,f'{gl(c0)}{R_GRP}',t,f'{gl(c0)}{R_GRP}:{gl(c1)}{R_GRP}')
        ws[f'{gl(c0)}{R_GRP}'].alignment = Alignment(horizontal='center',vertical='center')
    ws.row_dimensions[R_GRP].height = 20
    for c,h in ((1,'Категорія'),(2,'Назва на Booking'),(3,'Код')):
        col_head(ws,f'{gl(c)}{R_HEAD}',h)
    for c in range(WD0, WE1+1):
        if c == SPACER: continue
        n = (c - WD0 + 1) if c <= WD1 else (c - WE0 + 1)
        col_head(ws,f'{gl(c)}{R_HEAD}',n); ws[f'{gl(c)}{R_HEAD}'].number_format = '0'
    ws.row_dimensions[R_HEAD].height = 20
    for w,c in zip([30,30,10],'ABC'): ws.column_dimensions[c].width = w
    for c in range(WD0, WE1+1): ws.column_dimensions[gl(c)].width = 12.5
    ws.column_dimensions[gl(SPACER)].width = 2.6

    g = {3:'=IFERROR(MATCH($B$3,ПЛ_РЯДОК,0),0)',
         4:'=IFERROR(SUMIFS(ТАРИФИ_ЗНИЖКИ,ТАРИФИ_НАЗВИ,$B$4),0)',
         5:'=IF($B$5="BBSPA",ЗНИЖКА_BBSPA,$U$4)',
         6:'=IF($U$3=0,0,IFERROR(INDEX(ПЕРЕМИКАЧІ,IF($B$5="BBSPA",2,1),$U$3),0))',
         7:'=IF($U$3=0,0,IFERROR(INDEX(ПЕРЕМИКАЧІ,3,$U$3),0))',
         8:'=IF(OR($B$5="BB",$B$5="BBSPA",$B$5="ROSPABB"),1,0)',
         9:'=IF(OR($B$5="ROSPA",$B$5="ROSPABB"),1,0)',
        10:'=IF($U$8=0,"",IF($B$5="ROSPABB","BB",$B$5))',
        12:'=IFERROR(MATCH($B$3,ПРАЙС_ЛИСТИ,0),0)',
        11:('=IF($U$12=0,"період не заданий",IFERROR(IF(OR(INDEX(ПЛ_ДАТА_З,$U$12)="",INDEX(ПЛ_ДАТА_ПО,$U$12)=""),'
            '"період не заданий",'
            'TEXT(DAY(INDEX(ПЛ_ДАТА_З,$U$12)),"00")&"."&TEXT(MONTH(INDEX(ПЛ_ДАТА_З,$U$12)),"00")&"."&YEAR(INDEX(ПЛ_ДАТА_З,$U$12))'
            '&" – "&TEXT(DAY(INDEX(ПЛ_ДАТА_ПО,$U$12)),"00")&"."&TEXT(MONTH(INDEX(ПЛ_ДАТА_ПО,$U$12)),"00")&"."'
            '&YEAR(INDEX(ПЛ_ДАТА_ПО,$U$12))),"період не заданий"))'),
        13:('=IF($U$3=0,0,SUMPRODUCT(--(INDEX(БАЗА_ЦІНИ,0,$U$3)<>"")))')}
    labels = {3:'колонка прайс-листа',4:'знижка тарифу',5:'знижка на проживання',6:'перемикач: пакет',
              7:'перемикач: SPA',8:'потрібен пакет',9:'потрібен SPA',10:'ключ пакета',
              11:'період (текст)',12:'рядок прайс-листа',13:'заповнених баз'}
    HL, VL = gl(HELP0), gl(HELP0+1)
    for r,fml in g.items():
        ws[f'{HL}{r}'] = fml; ws[f'{HL}{r}'].font = f(9,color='6B7771')
        ws[f'{VL}{r}'] = labels[r]; ws[f'{VL}{r}'].font = f(9,it=True,color='9AA6A0')
    ws[f'{HL}1'] = 'ТЕХНІЧНІ ДАНІ — НЕ РЕДАГУВАТИ'; ws[f'{HL}1'].font = f(10,True,'6B7771')

    _rows(ws)
    for rng,src in (('B3','=ПРАЙС_ЛИСТИ'),('B4','=ТАРИФИ_НАЗВИ'),('B5','=СПИСОК_ПАКЕТІВ')):
        dv = DataValidation(type='list',formula1=src,allow_blank=False,showErrorMessage=True,
                            errorTitle='Оберіть зі списку',
                            error='Значення має бути з випадаючого списку — інакше прайс не порахується.')
        ws.add_data_validation(dv); dv.add(rng)

    ws.freeze_panes = f'{gl(WD0)}{R0}'
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GREEN_D
    ws.print_area = f'A1:{gl(WE1)}{R1}'
    ws.print_title_rows = f'1:{R_HEAD}'
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.left.text = 'Phoenix Східниця'
    ws.oddFooter.center.text = '&A'; ws.oddFooter.right.text = 'Сторінка &P з &N'
    return ws


def _rows(ws):
    for i in range(N_CAT):
        r, dr, band = R0+i, 5+i, i % 2 == 1
        m = f'MATCH($A{r},КАТЕГОРІЇ,0)'
        ws[f'A{r}'] = f"=IF('Довідник категорій'!B{dr}=\"\",\"\",'Довідник категорій'!B{dr})"
        ws[f'B{r}'] = f'=IF($A{r}="","",IFERROR(INDEX(ДОВ_BOOKING,{m}),""))'
        ws[f'C{r}'] = f'=IF($A{r}="","",IFERROR(INDEX(ДОВ_КОД,{m}),""))'
        out(ws,f'A{r}',None,band=band); ws[f'A{r}'].font = f(b=True)
        out(ws,f'B{r}',None,band=band); out(ws,f'C{r}',None,band=band)
        ws[f'C{r}'].alignment = Alignment(horizontal='center')
        ws[f'{H["ro_b"]}{r}'] = f'=IF(OR($U$3=0,$A{r}=""),0,IFERROR(INDEX(БАЗА_ЦІНИ,MATCH($A{r},БАЗА_КАТ,0),$U$3),0))'
        ws[f'{H["wm"]}{r}']   = f'=IF(OR($U$3=0,$A{r}=""),0,IFERROR(INDEX(МНОЖ_ВИХ,MATCH($A{r},МНОЖ_КАТ,0),$U$3),0))'
        ws[f'{H["ro_v"]}{r}'] = f'={H["ro_b"]}{r}*{H["wm"]}{r}'
        ws[f'{H["maxg"]}{r}'] = f'=IF($A{r}="",0,IFERROR(INDEX(ДОВ_МАКС,{m}),0))'
        ws[f'{H["incl"]}{r}'] = f'=IF($A{r}="",0,IFERROR(INDEX(ДОВ_ВКЛ,{m}),0))'
        ws[f'{H["dopl"]}{r}'] = f'=IF($A{r}="",0,IFERROR(INDEX(ДОВ_ДОПЛ_ДОР,{m}),0))'
        ws[f'{H["pkg_b"]}{r}'] = ('=IF(OR($U$3=0,$U$8=0),0,IFERROR(INDEX(ПАКЕТИ_ЦІНИ,'
                                  'MATCH($U$10&"|Будні",ПАКЕТИ_КЛЮЧ,0),$U$3),0))')
        ws[f'{H["pkg_v"]}{r}'] = ('=IF(OR($U$3=0,$U$8=0),0,IFERROR(INDEX(ПАКЕТИ_ЦІНИ,'
                                  'MATCH($U$10&"|Вихідні",ПАКЕТИ_КЛЮЧ,0),$U$3),0))')
        ws[f'{H["spa_b"]}{r}'] = (f'=IF(OR($U$3=0,$U$9=0,$A{r}=""),0,IFERROR(INDEX(SPA_ЦІНИ,'
                                  f'MATCH(INDEX(ДОВ_SPA,{m}),SPA_РІВЕНЬ,0),$U$3),0))')
        ws[f'{H["spa_v"]}{r}'] = (f'=IF(OR($U$3=0,$U$9=0,$A{r}=""),0,IFERROR(INDEX(SPA_ЦІНИ,'
                                  f'MATCH(INDEX(ДОВ_SPA,{m}),SPA_РІВЕНЬ,0),$U$3+1),0))')
        ws[f'{H["ok_b"]}{r}'] = (f'=IF(OR($A{r}="",{H["ro_b"]}{r}=0,AND($U$8=1,{H["pkg_b"]}{r}=0),'
                                 f'AND($U$9=1,{H["spa_b"]}{r}=0)),0,1)')
        ws[f'{H["ok_v"]}{r}'] = (f'=IF(OR($A{r}="",{H["ro_b"]}{r}=0,{H["wm"]}{r}=0,AND($U$8=1,{H["pkg_v"]}{r}=0),'
                                 f'AND($U$9=1,{H["spa_v"]}{r}=0)),0,1)')
        for k in H.values(): ws[f'{k}{r}'].font = f(9,color='6B7771')

        for c in range(WD0, WE1+1):
            if c == SPACER: continue
            wknd = c >= WE0
            n   = f'{gl(c)}${R_HEAD}'
            ok  = H['ok_v'] if wknd else H['ok_b']
            base= H['ro_v'] if wknd else H['ro_b']
            pkg = H['pkg_v'] if wknd else H['pkg_b']
            spa = H['spa_v'] if wknd else H['spa_b']
            ws.cell(row=r,column=c).value = (
                f'=IF(OR(${ok}{r}=0,{n}>${H["maxg"]}{r}),"",'
                f'ROUND(${base}{r}*(1-$U$5)/ОКРУГЛЕННЯ,0)*ОКРУГЛЕННЯ'
                f'+MAX(0,{n}-${H["incl"]}{r})*${H["dopl"]}{r}*(1-$U$5)'
                f'+{n}*${pkg}{r}*(1-$U$5*$U$6)'
                f'+${spa}{r}*(1-$U$5*$U$7))')
            out(ws,f'{gl(c)}{r}',None,MONEY,band=band)
            ws.cell(row=r,column=c).alignment = Alignment(horizontal='right')

    for r in range(R0, R1+1):
        edge(ws,f'C{r}',right=True); edge(ws,f'{gl(WD1)}{r}',right=True); edge(ws,f'{gl(WE0)}{r}',left=True)
    edge(ws,f'C{R_HEAD}',right=True); edge(ws,f'{gl(WD1)}{R_HEAD}',right=True); edge(ws,f'{gl(WE0)}{R_HEAD}',left=True)
    for c in range(HELP0, HELP0+14):
        ws.column_dimensions[gl(c)].hidden = True
