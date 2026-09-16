"""Дизайн-токени, успадковані з оригінального файлу Phoenix Bukovel."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection

FONT = "Calibri"           # оригінал: Carlito (метрично ідентичний, але немає у Windows Excel)

GREEN_D   = "1F4E3D"       # заголовок листа
GREEN_M   = "2F6B55"       # підзаголовок блоку
GREEN_PL  = "EAF3F0"       # пояснення / банер
GREEN_BAND= "F4F8F6"       # смугування рядків
INPUT     = "FFF2CC"       # ввід користувача
INPUT_ALT = "FFF9E6"       # ввід другого рівня
GREY_T    = "EDEFEE"       # технічна зона
RED_SOFT  = "FCE8E4"       # базові категорії
RED_INK   = "A8321E"

THIN = Side(style="thin", color="C9D4CD")
MED  = Side(style="medium", color="7E9488")     # межа блоку / групи
BOX  = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOX_L = Border(left=MED,  right=THIN, top=THIN, bottom=THIN)     # початок блоку прайс-листа
BOX_B = Border(left=THIN, right=THIN, top=THIN, bottom=MED)      # кінець групи категорій
BOX_LB= Border(left=MED,  right=THIN, top=THIN, bottom=MED)
BOX_R = Border(left=THIN, right=MED,  top=THIN, bottom=THIN)     # кінець описових колонок
BOX_RB= Border(left=THIN, right=MED,  top=THIN, bottom=MED)

def edge(ws, cell, left=False, bottom=False, right=False):
    """Додає товсту межу, зберігаючи наявні тонкі."""
    b = ws[cell].border
    ws[cell].border = Border(
        left  = MED if left   else (b.left   or THIN),
        right = MED if right  else (b.right  or THIN),
        top   = b.top or THIN,
        bottom= MED if bottom else (b.bottom or THIN))

def gutter(ws, row, c0, c1, height=8):
    """Порожній рядок-роздільник між блоками."""
    ws.row_dimensions[row].height = height
    for c in range(c0, c1 + 1):
        ws.cell(row=row, column=c).fill = fill(GREEN_BAND)

def fill(hexcolor):
    return PatternFill("solid", fgColor=hexcolor)

def f(sz=11, b=False, color="000000", it=False):
    return Font(name=FONT, size=sz, bold=b, color=color, italic=it)

def title(ws, cell, text, span, sz=14):
    ws[cell] = text
    ws[cell].font = f(sz, True, "FFFFFF")
    ws[cell].fill = fill(GREEN_D)
    ws[cell].alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.merge_cells(span)
    ws.row_dimensions[ws[cell].row].height = 27.75

def banner(ws, cell, text, span, h=30):
    ws[cell] = text
    ws[cell].font = f(10, color="2A3B33", it=True)
    ws[cell].fill = fill(GREEN_PL)
    ws[cell].alignment = Alignment(vertical="center", wrap_text=True, indent=1)
    ws.merge_cells(span)
    ws.row_dimensions[ws[cell].row].height = h

def block_head(ws, cell, text, span=None):
    ws[cell] = text
    ws[cell].font = f(11, True, "FFFFFF")
    ws[cell].fill = fill(GREEN_M)
    ws[cell].alignment = Alignment(vertical="center", indent=1)
    if span:
        ws.merge_cells(span)

def col_head(ws, cell, text, wrap=True):
    ws[cell] = text
    ws[cell].font = f(11, True, "FFFFFF")
    ws[cell].fill = fill(GREEN_D)
    ws[cell].alignment = Alignment(vertical="center", horizontal="center", wrap_text=wrap)
    ws[cell].border = BOX

def inp(ws, cell, value=None, numfmt=None, alt=False):
    if value is not None:
        ws[cell] = value
    ws[cell].fill = fill(INPUT_ALT if alt else INPUT)
    ws[cell].font = f()
    ws[cell].border = BOX
    ws[cell].protection = Protection(locked=False)   # жовті клітинки лишаються редагованими
    if numfmt:
        ws[cell].number_format = numfmt

def out(ws, cell, value=None, numfmt=None, bold=False, band=False):
    if value is not None:
        ws[cell] = value
    ws[cell].font = f(b=bold)
    ws[cell].border = BOX
    if band:
        ws[cell].fill = fill(GREEN_BAND)
    if numfmt:
        ws[cell].number_format = numfmt

def label(ws, cell, text, bold=False, wrap=False):
    ws[cell] = text
    ws[cell].font = f(b=bold)
    ws[cell].alignment = Alignment(vertical="center", wrap_text=wrap)

MONEY = "#,##0"
MULT  = "0.0000"
PCT   = "0.0%"
DATE  = "DD.MM.YYYY"


def protect(ws):
    """Захист без пароля: редагувати можна лише жовті клітинки."""
    ws.protection.sheet = True
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False
    ws.protection.formatCells = True
    ws.protection.formatColumns = True
    ws.protection.formatRows = True
    ws.protection.sort = True
    ws.protection.autoFilter = True
