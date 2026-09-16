from openpyxl.styles import Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.utils import get_column_letter as gl
from style import *
from names_s import *

# ═════════ КОНТРОЛЬ І ПОРІВНЯННЯ ═════════
KR0 = 7; KR1 = KR0 + N_CAT - 1
HK = ['ro_b','wm','ro_v','maxg','incl','ming','dopl_a','dopl_k','pa_b','pk_b','pa_v','pk_v','spa_b','spa_v','okd','oke']
H1_0, H2_0 = 16, 34
NTOT = '($D$4+$F$4)'
def hc(start): return {k: gl(start+i) for i,k in enumerate(HK)}

def build_control(wb, data):
    ws = wb.create_sheet('Контроль і порівняння')
    title(ws,'A1','ПОРІВНЯННЯ ТА КОНТРОЛЬ ПРАЙС-ЛИСТІВ — PHOENIX СХІДНИЦЯ','A1:I1')
    for w,c in zip([30,13,13,13,13,12,12,11,11],'ABCDEFGHI'): ws.column_dimensions[c].width = w
    for addr,lab,val,fmt in (('B3','Прайс-лист 1','0сінь 26 0-10%',None),('D3','Тариф 1','Rack Rate',None),
                             ('F3','Прайс-лист 2','0сінь 26 20-30%',None),('H3','Тариф 2','Rack Rate',None),
                             ('B4','Пакет','RO',None),('D4','Дорослих',2,'0'),('F4','Дітей',0,'0')):
        label(ws, gl(ws[addr].column-1)+addr[1:], lab, bold=True); inp(ws,addr,val,fmt)
    for rng,src in (('B3','=ПРАЙС_ЛИСТИ'),('F3','=ПРАЙС_ЛИСТИ'),('D3','=ТАРИФИ_НАЗВИ'),
                    ('H3','=ТАРИФИ_НАЗВИ'),('B4','=СПИСОК_ПАКЕТІВ')):
        dv = DataValidation(type='list',formula1=src,allow_blank=False); ws.add_data_validation(dv); dv.add(rng)
    dvg = DataValidation(type='whole',operator='between',formula1='0',formula2='30',allow_blank=False,
                         showErrorMessage=True,errorTitle='Кількість гостей',error='Ціле число від 0 до 30.')
    ws.add_data_validation(dvg); dvg.add('D4'); dvg.add('F4')
    banner(ws,'A5','Порівняння показує ціну за одну ніч без тварин — для буднього і вихідного дня одночасно. '
                   'Значення поза місткістю категорії не показуються.','A5:I5',28)
    for i,h in enumerate(['Категорія','Прайс 1\nбудні','Прайс 1\nвихідні','Прайс 2\nбудні','Прайс 2\nвихідні',
                          'Δ грн\nбудні','Δ грн\nвихідні','Δ %\nбудні','Δ %\nвихідні']):
        col_head(ws,f'{gl(i+1)}6',h)
    ws.row_dimensions[6].height = 34
    _kglob(ws); _khelp(ws,'$B$3','$D$3',H1_0); _khelp(ws,'$F$3','$H$3',H2_0)
    for i in range(N_CAT):
        r, dr, band = KR0+i, 5+i, i % 2 == 1
        ws[f'A{r}'] = f"=IF('Довідник категорій'!B{dr}=\"\",\"\",'Довідник категорій'!B{dr})"
        out(ws,f'A{r}',None,band=band); ws[f'A{r}'].font = f(b=True)
        ws[f'B{r}'] = _kprice(H1_0,False,r); ws[f'C{r}'] = _kprice(H1_0,True,r)
        ws[f'D{r}'] = _kprice(H2_0,False,r); ws[f'E{r}'] = _kprice(H2_0,True,r)
        ws[f'F{r}'] = f'=IF(OR(B{r}="",D{r}=""),"",D{r}-B{r})'
        ws[f'G{r}'] = f'=IF(OR(C{r}="",E{r}=""),"",E{r}-C{r})'
        ws[f'H{r}'] = f'=IF(OR(B{r}="",D{r}="",B{r}=0),"",D{r}/B{r}-1)'
        ws[f'I{r}'] = f'=IF(OR(C{r}="",E{r}="",C{r}=0),"",E{r}/C{r}-1)'
        for c in 'BCDEFG': out(ws,f'{c}{r}',None,MONEY,band=band)
        for c in 'HI':     out(ws,f'{c}{r}',None,PCT,band=band)
    for rng in (f'F{KR0}:G{KR1}', f'H{KR0}:I{KR1}'):
        ws.conditional_formatting.add(rng, CellIsRule(operator='lessThan',formula=['0'],font=f(11,color=RED_INK)))
    _kcheck(ws)
    for c in range(14, H2_0+len(HK)): ws.column_dimensions[gl(c)].hidden = True
    ws.freeze_panes = f'B{KR0}'; ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GREEN_M
    ws.print_area = f'A1:I{KR1}'; ws.print_title_rows = '1:6'
    ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws

def _kglob(ws):
    ws['N1'] = 'ТЕХНІЧНІ ДАНІ'; ws['N1'].font = f(10,True,'6B7771')
    for r,fml in ((3,"=IFERROR(MATCH('Прайс'!$B$3,ПЛ_РЯДОК,0),0)"),
                  (5,'=IF(OR($B$4="BB",$B$4="BBSPA",$B$4="ROSPABB"),1,0)'),
                  (6,'=IF($N$5=0,"",IF($B$4="ROSPABB","BB",$B$4))'),
                  (7,'=IF(OR($B$4="ROSPA",$B$4="ROSPABB"),1,0)')):
        ws[f'N{r}'] = fml; ws[f'N{r}'].font = f(9,color='6B7771')

def _khelp(ws, pl, tar, start):
    c = hc(start); H = gl(start)
    ws[f'{H}2'] = f'=IFERROR(MATCH({pl},ПЛ_РЯДОК,0),0)'
    ws[f'{H}3'] = f'=IF($B$4="BBSPA",ЗНИЖКА_BBSPA,IFERROR(SUMIFS(ТАРИФИ_ЗНИЖКИ,ТАРИФИ_НАЗВИ,{tar}),0))'
    ws[f'{H}4'] = f'=IF({H}$2=0,0,IFERROR(INDEX(ПЕРЕМИКАЧІ,IF($B$4="BBSPA",2,1),{H}$2),0))'
    ws[f'{H}5'] = f'=IF({H}$2=0,0,IFERROR(INDEX(ПЕРЕМИКАЧІ,3,{H}$2),0))'
    for r in range(2,6): ws[f'{H}{r}'].font = f(9,color='6B7771')
    for i in range(N_CAT):
        r = KR0+i; m = f'MATCH($A{r},КАТЕГОРІЇ,0)'
        ws[f'{c["ro_b"]}{r}'] = f'=IF(OR({H}$2=0,$A{r}=""),0,IFERROR(INDEX(БАЗА_ЦІНИ,MATCH($A{r},БАЗА_КАТ,0),{H}$2),0))'
        ws[f'{c["wm"]}{r}']   = f'=IF(OR({H}$2=0,$A{r}=""),0,IFERROR(INDEX(МНОЖ_ВИХ,MATCH($A{r},МНОЖ_КАТ,0),{H}$2),0))'
        ws[f'{c["ro_v"]}{r}'] = f'={c["ro_b"]}{r}*{c["wm"]}{r}'
        ws[f'{c["maxg"]}{r}'] = f'=IF($A{r}="",0,IFERROR(INDEX(ДОВ_МАКС,{m}),0))'
        ws[f'{c["incl"]}{r}'] = f'=IF($A{r}="",0,IFERROR(INDEX(ДОВ_ВКЛ,{m}),0))'
        ws[f'{c["ming"]}{r}'] = f'=IF($A{r}="",0,IFERROR(INDEX(ДОВ_МІН,{m}),0))'
        ws[f'{c["dopl_a"]}{r}'] = f'=IF($A{r}="",0,IFERROR(INDEX(ДОВ_ДОПЛ_ДОР,{m}),0))'
        ws[f'{c["dopl_k"]}{r}'] = f'=IF($A{r}="",0,IFERROR(INDEX(ДОВ_ДОПЛ_ДІТ,{m}),0))'
        for key,day,off in (('pa_b','Будні',0),('pk_b','Будні',1),('pa_v','Вихідні',0),('pk_v','Вихідні',1)):
            ws[f'{c[key]}{r}'] = (f'=IF(OR({H}$2=0,$N$5=0),0,IFERROR(INDEX(ПАКЕТИ_ЦІНИ,'
                                  f'MATCH($N$6&"|{day}",ПАКЕТИ_КЛЮЧ,0),{H}$2+{off}),0))')
        for key,off in (('spa_b',0),('spa_v',1)):
            ws[f'{c[key]}{r}'] = (f'=IF(OR({H}$2=0,$N$7=0,$A{r}=""),0,IFERROR(INDEX(SPA_ЦІНИ,'
                                  f'MATCH(INDEX(ДОВ_SPA,{m}),SPA_РІВЕНЬ,0),{H}$2+{off}),0))')
        ws[f'{c["okd"]}{r}'] = (f'=IF(OR($A{r}="",{c["ro_b"]}{r}=0,AND($N$5=1,$D$4>0,{c["pa_b"]}{r}=0),'
                                f'AND($N$5=1,$F$4>0,{c["pk_b"]}{r}=0),AND($N$7=1,{c["spa_b"]}{r}=0)),0,1)')
        ws[f'{c["oke"]}{r}'] = (f'=IF(OR($A{r}="",{c["ro_b"]}{r}=0,{c["wm"]}{r}=0,AND($N$5=1,$D$4>0,{c["pa_v"]}{r}=0),'
                                f'AND($N$5=1,$F$4>0,{c["pk_v"]}{r}=0),AND($N$7=1,{c["spa_v"]}{r}=0)),0,1)')
        for k in c.values(): ws[f'{k}{r}'].font = f(9,color='6B7771')

def _kprice(start, wknd, r):
    c = hc(start); H = gl(start)
    ok  = c['oke'] if wknd else c['okd']
    ro  = c['ro_v'] if wknd else c['ro_b']
    pa  = c['pa_v'] if wknd else c['pa_b']
    pk  = c['pk_v'] if wknd else c['pk_b']
    spa = c['spa_v'] if wknd else c['spa_b']
    extra_a = f'MAX(0,$D$4-{c["incl"]}{r})'
    extra_k = f'(MAX(0,{NTOT}-{c["incl"]}{r})-{extra_a})'
    return (f'=IF(OR({ok}{r}=0,{NTOT}=0,{NTOT}>{c["maxg"]}{r},{NTOT}<{c["ming"]}{r}),"",'
            f'ROUND({ro}{r}*(1-{H}$3)/ОКРУГЛЕННЯ,0)*ОКРУГЛЕННЯ'
            f'+({extra_a}*{c["dopl_a"]}{r}+{extra_k}*{c["dopl_k"]}{r})*(1-{H}$3)'
            f'+($D$4*{pa}{r}+$F$4*{pk}{r})*(1-{H}$3*{H}$4)'
            f'+{spa}{r}*(1-{H}$3*{H}$5))')

def _kcheck(ws):
    block_head(ws,'K3','КОНТРОЛЬ АКТИВНОГО ПРАЙС-ЛИСТА','K3:M3')
    ws.column_dimensions['K'].width = 46; ws.column_dimensions['L'].width = 15; ws.column_dimensions['M'].width = 3
    checks = [(4,'Незаповнені базові RO-ціни (заповнених категорій)',
               f'=COUNTIF(КАТЕГОРІЇ,"<>")*2-COUNT(INDEX(БАЗА_ЦІНИ,0,$N$3))-COUNT(INDEX(БАЗА_ЦІНИ,0,$N$3+1))'),
              (5,'Незаповнені пакети BB',
               '=4-COUNT(INDEX(ПАКЕТИ_ЦІНИ,1,$N$3),INDEX(ПАКЕТИ_ЦІНИ,1,$N$3+1),'
               'INDEX(ПАКЕТИ_ЦІНИ,2,$N$3),INDEX(ПАКЕТИ_ЦІНИ,2,$N$3+1))'),
              (6,'Дублікати назв прайс-листів',
               '=SUMPRODUCT((ПРАЙС_ЛИСТИ<>"")*(COUNTIF(ПРАЙС_ЛИСТИ,ПРАЙС_ЛИСТИ&"")>1))')]
    for r,lab,fml in checks:
        label(ws,f'K{r}',lab); ws[f'L{r}'] = fml; out(ws,f'L{r}',None,'0')
        ws.conditional_formatting.add(f'L{r}', CellIsRule(operator='greaterThan',formula=['0'],
                                     fill=fill(RED_SOFT), font=f(11,True,RED_INK)))
    ws['K8'] = '=IF(SUM($L$4:$L$6)=0,"ГОТОВО ДО ПУБЛІКАЦІЇ","ПОТРІБНЕ ЗАПОВНЕННЯ")'
    ws['K8'].font = f(12,True); ws['K8'].alignment = Alignment(horizontal='center',vertical='center')
    ws.merge_cells('K8:L8'); ws.row_dimensions[8].height = 26
    ws.conditional_formatting.add('K8:L8', FormulaRule(formula=['$K$8="ГОТОВО ДО ПУБЛІКАЦІЇ"'],
                                  fill=fill('E1EFE7'), font=f(12,True,'2E6B4F'), stopIfTrue=True))
    ws.conditional_formatting.add('K8:L8', FormulaRule(formula=['$K$8="ПОТРІБНЕ ЗАПОВНЕННЯ"'],
                                  fill=fill(RED_SOFT), font=f(12,True,RED_INK), stopIfTrue=True))
    info = [(10,'Пакет BBSPA заповнено?',
             '=IF(COUNT(INDEX(ПАКЕТИ_ЦІНИ,3,$N$3),INDEX(ПАКЕТИ_ЦІНИ,3,$N$3+1),'
             'INDEX(ПАКЕТИ_ЦІНИ,4,$N$3),INDEX(ПАКЕТИ_ЦІНИ,4,$N$3+1))=4,"так","ні — пакет недоступний")',None),
            (11,'Фіксований SPA заповнено? (ROSPA / ROSPABB)',
             '=IF(COUNT(INDEX(SPA_ЦІНИ,1,$N$3),INDEX(SPA_ЦІНИ,1,$N$3+1),'
             'INDEX(SPA_ЦІНИ,2,$N$3),INDEX(SPA_ЦІНИ,2,$N$3+1))=4,"так","ні — пакети недоступні")',None),
            (12,'Прайс-листів у моделі','=COUNTIF(ПРАЙС_ЛИСТИ,"<>")','0'),
            (13,'Категорій у моделі','=COUNTIF(КАТЕГОРІЇ,"<>")','0'),
            (14,'Крок округлення, грн','=ОКРУГЛЕННЯ',MONEY)]
    for r,lab,fml,fmt in info:
        label(ws,f'K{r}',lab); ws[f'L{r}'] = fml; out(ws,f'L{r}',None,fmt,band=True)
    label(ws,'K16','BBSPA і фіксований SPA не входять до чек-листа: пакети можуть бути свідомо не задіяні у сезоні.')
    ws['K16'].font = f(9,it=True,color='6B7771')
    ws['K16'].alignment = Alignment(wrap_text=True,vertical='top'); ws.merge_cells('K16:L17')


# ═════════ МОДЕЛЮВАННЯ ═════════
def build_model(wb, data):
    ws = wb.create_sheet('Моделювання')
    R0 = 8; R1 = R0 + N_CAT - 1
    title(ws,'A1','МОДЕЛЮВАННЯ ЦІН — ДВА СЦЕНАРІЇ У ПОРІВНЯННІ','A1:Q1')
    label(ws,'A3','Прайс-лист-основа',bold=True); inp(ws,'B3','0сінь 26 0-10%')
    dv = DataValidation(type='list',formula1='=ПРАЙС_ЛИСТИ',allow_blank=False); ws.add_data_validation(dv); dv.add('B3')
    banner(ws,'A4','Жовті колонки — ввід ціни буднього дня та множника вихідних. Сценарії нічого не змінюють у діючому прайсі: щоб застосувати сценарій, '
                   'скопіюйте його ціни у блок відповідного прайс-листа на «Керування прайсом». '
                   'Ціни показані за базовою місткістю категорії, без пакетів.','A4:Q4',34)
    GR = [(3,4,'ПОТОЧНО'),(5,8,'СЦЕНАРІЙ 1'),(9,12,'СЦЕНАРІЙ 2'),(13,14,'СЦЕНАРІЙ 2 vs 1')]
    for c0,c1,t in GR:
        block_head(ws,f'{gl(c0)}6',t,f'{gl(c0)}6:{gl(c1)}6')
        ws[f'{gl(c0)}6'].alignment = Alignment(horizontal='center',vertical='center')
    HEAD = ['Категорія','Код','Будні, грн','Множник вих.',
            'Будні, грн','Множник вих.','Δ грн буд.','Δ % буд.',
            'Будні, грн','Множник вих.','Δ грн буд.','Δ % буд.','Δ грн буд.','Δ % буд.']
    for i,h in enumerate(HEAD): col_head(ws,f'{gl(i+1)}7',h)
    ws.row_dimensions[7].height = 30
    for w,c in zip([30,10],'AB'): ws.column_dimensions[c].width = w
    for c in range(3,15): ws.column_dimensions[gl(c)].width = 13
    ws['S2'] = '=IFERROR(MATCH($B$3,ПЛ_РЯДОК,0),0)'; ws['S2'].font = f(9,color='6B7771')
    ws['S1'] = 'ТЕХНІЧНІ ДАНІ'; ws['S1'].font = f(10,True,'6B7771')
    for c in ('S','T'): ws.column_dimensions[c].hidden = True
    for i in range(N_CAT):
        r, dr, band = R0+i, 5+i, i % 2 == 1
        m = f'MATCH($A{r},КАТЕГОРІЇ,0)'
        ws[f'A{r}'] = f"=IF('Довідник категорій'!B{dr}=\"\",\"\",'Довідник категорій'!B{dr})"
        ws[f'B{r}'] = f'=IF($A{r}="","",IFERROR(INDEX(ДОВ_КОД,{m}),""))'
        ws[f'C{r}'] = f'=IF(OR($S$2=0,$A{r}=""),"",IFERROR(INDEX(БАЗА_ЦІНИ,MATCH($A{r},БАЗА_КАТ,0),$S$2),""))'
        ws[f'D{r}'] = f'=IF(OR($S$2=0,$A{r}=""),"",IFERROR(INDEX(МНОЖ_ВИХ,MATCH($A{r},МНОЖ_КАТ,0),$S$2),""))'
        out(ws,f'A{r}',None,band=band); ws[f'A{r}'].font = f(b=True)
        out(ws,f'B{r}',None,band=band); ws[f'B{r}'].alignment = Alignment(horizontal='center')
        out(ws,f'C{r}',None,MONEY,band=band); out(ws,f'D{r}',None,MULT,band=band)
        for base,alt in ((5,False),(9,True)):
            cb,cv,cd,cp = (gl(base+k) for k in range(4))
            inp(ws,f'{cb}{r}',None,MONEY,alt=alt); inp(ws,f'{cv}{r}',None,MULT,alt=alt)
            ws[f'{cd}{r}'] = f'=IF(OR({cb}{r}="",$C{r}=""),"",{cb}{r}-$C{r})'
            ws[f'{cp}{r}'] = f'=IF(OR({cb}{r}="",$C{r}="",$C{r}=0),"",{cb}{r}/$C{r}-1)'
            out(ws,f'{cd}{r}',None,MONEY,band=band); out(ws,f'{cp}{r}',None,PCT,band=band)
        ws[f'M{r}'] = f'=IF(OR(I{r}="",E{r}=""),"",I{r}-E{r})'
        ws[f'N{r}'] = f'=IF(OR(I{r}="",E{r}="",E{r}=0),"",I{r}/E{r}-1)'
        out(ws,f'M{r}',None,MONEY,band=band); out(ws,f'N{r}',None,PCT,band=band)
    for rng in (f'G{R0}:G{R1}',f'H{R0}:H{R1}',f'K{R0}:K{R1}',f'L{R0}:L{R1}',f'M{R0}:M{R1}',f'N{R0}:N{R1}'):
        ws.conditional_formatting.add(rng, CellIsRule(operator='lessThan',formula=['0'],font=f(11,color=RED_INK)))
    dvp = DataValidation(type='decimal',operator='greaterThanOrEqual',formula1='0',allow_blank=True,
                         showErrorMessage=True,errorTitle='Ціна',error='Ціна не може бути відʼємною.')
    ws.add_data_validation(dvp); dvp.add(f'E{R0}:F{R1}'); dvp.add(f'I{R0}:J{R1}')
    ws.freeze_panes = f'C{R0}'; ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = '8A7128'
    ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


# ═════════ БАЗОВІ ЦІНИ ═════════
def build_base(wb, data):
    ws = wb.create_sheet('Базові ціни')
    R0 = 7; R1 = R0 + N_CAT - 1
    title(ws,'A1','БАЗОВІ ЦІНИ — ПЕРЕГЛЯД АКТИВНОГО ПРАЙСУ','A1:H1')
    label(ws,'A3','Активний прайс-лист',bold=True)
    ws['B3'] = "='Прайс'!$B$3"; ws['B3'].fill = fill(GREEN_PL); ws['B3'].font = f(b=True,color=GREEN_D)
    banner(ws,'A4','Усі категорії Східниці — базові: кожна має власну RO-ціну буднього дня. '
                   'Ціна вихідних рахується множником. Редагування — на «Керування прайсом».','A4:H4',26)
    for i,h in enumerate(['Категорія','Код','RO будні, грн','Множник вихідних','RO вихідні, грн',
                          'Δ буд.→вих., грн','Включено у базову','Макс. гостей']):
        col_head(ws,f'{gl(i+1)}6',h)
    ws.row_dimensions[6].height = 30
    for w,c in zip([30,10,15,16,16,17,17,14],'ABCDEFGH'): ws.column_dimensions[c].width = w
    ws['I2'] = "=IFERROR(MATCH('Прайс'!$B$3,ПЛ_РЯДОК,0),0)"; ws['I2'].font = f(9,color='6B7771')
    ws.column_dimensions['I'].hidden = True
    for i in range(N_CAT):
        r, dr, band = R0+i, 5+i, i % 2 == 1
        m = f'MATCH($A{r},КАТЕГОРІЇ,0)'
        ws[f'A{r}'] = f"=IF('Довідник категорій'!B{dr}=\"\",\"\",'Довідник категорій'!B{dr})"
        ws[f'B{r}'] = f'=IF($A{r}="","",IFERROR(INDEX(ДОВ_КОД,{m}),""))'
        ws[f'C{r}'] = f'=IF(OR($I$2=0,$A{r}=""),"",IFERROR(INDEX(БАЗА_ЦІНИ,MATCH($A{r},БАЗА_КАТ,0),$I$2),""))'
        ws[f'D{r}'] = f'=IF(OR($I$2=0,$A{r}=""),"",IFERROR(INDEX(МНОЖ_ВИХ,MATCH($A{r},МНОЖ_КАТ,0),$I$2),""))'
        ws[f'E{r}'] = f'=IF(OR($C{r}="",$D{r}=""),"",ROUND($C{r}*$D{r}/ОКРУГЛЕННЯ,0)*ОКРУГЛЕННЯ)'
        ws[f'F{r}'] = f'=IF(OR($C{r}="",$E{r}=""),"",$E{r}-$C{r})'
        ws[f'G{r}'] = f'=IF($A{r}="","",IFERROR(INDEX(ДОВ_ВКЛ,{m}),""))'
        ws[f'H{r}'] = f'=IF($A{r}="","",IFERROR(INDEX(ДОВ_МАКС,{m}),""))'
        out(ws,f'A{r}',None,band=band); ws[f'A{r}'].font = f(b=True)
        for c,fmt in (('B',None),('C',MONEY),('D',MULT),('E',MONEY),('F',MONEY),('G','0'),('H','0')):
            out(ws,f'{c}{r}',None,fmt,band=band)
            if c in ('B','G','H'): ws[f'{c}{r}'].alignment = Alignment(horizontal='center')
    ws.freeze_panes = f'C{R0}'; ws.sheet_view.showGridLines = False
    return ws


# ═════════ ЖУРНАЛ ЗМІН ═════════
def build_log(wb, data, today):
    ws = wb.create_sheet('Журнал змін')
    title(ws,'A1','ЖУРНАЛ ЗМІН','A1:F1')
    banner(ws,'A2','Заповнюється вручну. Правило: діючий прайс-лист не перезаписуємо — переводимо у статус «Архів» '
                   'і створюємо новий рядок на «Керування прайсом». Тоді порівняння рік/рік лишається можливим.','A2:F2',30)
    for i,h in enumerate(['Дата','Автор','Лист','Що змінено','Було','Стало']): col_head(ws,f'{gl(i+1)}4',h)
    for w,c in zip([13,20,26,46,20,20],'ABCDEF'): ws.column_dimensions[c].width = w
    rows = [(today,'Claude (перебудова)','Довідник категорій',
             'Premium Chalet: макс. гостей — узгоджено «7, 8 дорослий по додатковому місцю»','7','8'),
            (today,'Claude (перебудова)','Керування прайсом',
             'Додано пакети ROSPA / ROSPABB: блок фіксованого SPA-візиту (малий / великий обʼєкт × будні / вихідні)','—','порожній блок'),
            (today,'Claude (перебудова)','Керування прайсом',
             'Доплату за додаткове місце перенесено з глобальних параметрів у «Довідник категорій» (по категоріях)',
             'глобальна','по категорії'),
            (today,'Claude (перебудова)','Усі листи',
             'Структуру приведено до прайс-листа Буковеля v2.2; множників категорій немає — усі категорії базові','—','—')]
    for i,row in enumerate(rows):
        r = 5+i
        for j,v in enumerate(row):
            out(ws,f'{gl(j+1)}{r}',v, DATE if j==0 else None, band=(i%2==1))
            ws[f'{gl(j+1)}{r}'].alignment = Alignment(wrap_text=True,vertical='top')
    for r in range(5+len(rows), 60):
        for j in range(6): inp(ws,f'{gl(j+1)}{r}',None, DATE if j==0 else None)
    ws.freeze_panes = 'A5'; ws.sheet_view.showGridLines = False
    return ws


# ═════════ ДОДАТКОВІ ПОСЛУГИ ═════════
def build_extra(wb, data):
    ws = wb.create_sheet('Додаткові послуги')
    title(ws,'A1','ДОДАТКОВІ ПОСЛУГИ — ІНФОРМАЦІЙНО','A1:E1')
    banner(ws,'A2','Цей лист НЕ бере участі в розрахунку прайсу. Ціни сніданків і SPA, які реально рахує система, '
                   'задаються на «Керування прайсом» у блоках «ПАКЕТИ BB / BBSPA» та «ФІКСОВАНИЙ SPA-ВІЗИТ». '
                   'Тут — довідкові умови продажу та взаєморозрахунків.','A2:E2',40)
    for w,c in zip([32,20,20,16,26],'ABCDE'): ws.column_dimensions[c].width = w
    r = 4
    for row in data['dop']:
        if all(v in (None,'') for v in row): r += 1; continue
        first = str(row[0] or '')
        for j,v in enumerate(row):
            if v is None: continue
            cell = f'{gl(j+1)}{r}'; ws[cell] = v; ws[cell].font = f()
            if isinstance(v,(int,float)) and j > 0: ws[cell].number_format = MONEY
        if first.startswith('Додаткові послуги'):
            block_head(ws,f'A{r}',first,f'A{r}:E{r}')
        elif first.strip() == 'Назва послуги':
            for j in range(5):
                if row[j] is not None: col_head(ws,f'{gl(j+1)}{r}',row[j])
        r += 1
    ws.sheet_view.showGridLines = False
    return ws
